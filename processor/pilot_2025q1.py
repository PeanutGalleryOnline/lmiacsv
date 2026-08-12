"""Normalize the official English 2025 Q1 LMIA employer-list workbooks.

The output is intentionally narrow: one canonical CSV for audit and compact JSON
aggregates for the static publication. Each record retains enough provenance to
resolve it back to an official workbook row.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import unicodedata
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass, fields
from pathlib import Path
from typing import Any, Iterable


WORKSPACE_ROOT = Path(__file__).resolve().parents[1]
MANIFEST_PATH = WORKSPACE_ROOT / "content" / "releases" / "2025q1.json"

PROVINCE_CODES = {
    "alberta": "AB",
    "british columbia": "BC",
    "manitoba": "MB",
    "new brunswick": "NB",
    "newfoundland and labrador": "NL",
    "nova scotia": "NS",
    "northwest territories": "NT",
    "nunavut": "NU",
    "ontario": "ON",
    "prince edward island": "PE",
    "quebec": "QC",
    "saskatchewan": "SK",
    "yukon": "YT",
    "employers carrying on business in canada with head office outside of canada": "INT",
}

POSTAL_RE = re.compile(
    r"\b([ABCEGHJ-NPRSTVXY]\d[ABCEGHJ-NPRSTVWXYZ])\s?(\d[ABCEGHJ-NPRSTVWXYZ]\d)\b",
    re.I,
)
OCCUPATION_RE = re.compile(r"^(\d{5})\s*[-–]\s*(.+)$")


@dataclass(frozen=True)
class CanonicalRecord:
    record_id: str
    release_id: str
    period: str
    decision_status: str
    language: str
    province_raw: str
    province_code: str
    program_stream_raw: str
    program_stream: str
    employer_raw: str
    employer: str
    employer_key: str
    address_raw: str
    address: str
    postal_code: str
    occupation_raw: str
    noc_version: str
    noc_code: str
    occupation_title: str
    occupation_label: str
    incorporate_status_raw: str
    incorporate_status: str
    lmia_measure: str
    lmia_count: int
    position_measure: str
    position_count: int
    source_dataset_id: str
    source_resource_id: str
    source_resource_url: str
    source_download_url: str
    source_file: str
    source_sheet: str
    source_row: int
    source_sha256: str


def load_manifest(path: Path = MANIFEST_PATH) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(unicodedata.normalize("NFKC", str(value)).strip().split())


def raw_text(value: Any) -> str:
    if value is None:
        return ""
    return unicodedata.normalize("NFC", str(value))


def search_key(value: str) -> str:
    folded = unicodedata.normalize("NFKD", value.casefold())
    letters = "".join(
        char if char.isalnum() else " "
        for char in folded
        if not unicodedata.combining(char)
    )
    return " ".join(letters.split())


def integer(value: Any, label: str, row_number: int, minimum: int = 1) -> int:
    if isinstance(value, bool):
        raise ValueError(f"Row {row_number}: {label} cannot be boolean")
    if isinstance(value, int):
        parsed = value
    elif isinstance(value, float) and value.is_integer():
        parsed = int(value)
    else:
        text = clean_text(value)
        if not text.isdigit():
            raise ValueError(f"Row {row_number}: {label} is not an integer: {value!r}")
        parsed = int(text)
    if parsed < minimum:
        raise ValueError(
            f"Row {row_number}: {label} must be at least {minimum}: {parsed}"
        )
    return parsed


def parse_occupation(value: str, row_number: int) -> tuple[str, str]:
    match = OCCUPATION_RE.match(value)
    if not match:
        raise ValueError(f"Row {row_number}: cannot parse NOC 2021 occupation: {value!r}")
    return match.group(1), match.group(2).strip()


def postal_code(address: str) -> str:
    match = POSTAL_RE.search(address)
    return f"{match.group(1).upper()} {match.group(2).upper()}" if match else ""


def stable_record_id(resource_id: str, sheet: str, row_number: int) -> str:
    identity = f"{resource_id}:{sheet}:{row_number}".encode()
    return f"lmia-{hashlib.sha256(identity).hexdigest()[:20]}"


def normalize_row(
    resource: dict[str, Any], row_number: int, values: Iterable[Any]
) -> CanonicalRecord:
    row = list(values)
    if len(row) != 8:
        raise ValueError(f"Row {row_number}: expected 8 cells, found {len(row)}")

    (
        province_raw,
        stream_raw,
        employer_raw,
        address_raw,
        occupation_raw,
        incorporate_raw,
        lmias_raw,
        positions_raw,
    ) = row
    province = clean_text(province_raw)
    stream = clean_text(stream_raw)
    employer = clean_text(employer_raw)
    address = clean_text(address_raw)
    occupation = clean_text(occupation_raw)
    incorporate_status = clean_text(incorporate_raw)

    required = {
        "province": province,
        "program stream": stream,
        "employer": employer,
        "address": address,
        "occupation": occupation,
    }
    missing = [label for label, value in required.items() if not value]
    if missing:
        raise ValueError(f"Row {row_number}: missing {', '.join(missing)}")

    province_code = PROVINCE_CODES.get(province.casefold())
    if province_code is None:
        raise ValueError(f"Row {row_number}: unresolved province: {province!r}")

    noc_code, occupation_title = parse_occupation(occupation, row_number)
    status = resource["status"]
    lmia_measure = "approved" if status == "positive" else "requested"
    position_measure = (
        "approved" if status == "positive" else "requested_on_negative_decision"
    )

    return CanonicalRecord(
        record_id=stable_record_id(resource["resource_id"], resource["sheet"], row_number),
        release_id=resource["id"],
        period="2025-Q1",
        decision_status=status,
        language=resource["language"],
        province_raw=raw_text(province_raw),
        province_code=province_code,
        program_stream_raw=raw_text(stream_raw),
        program_stream=stream,
        employer_raw=raw_text(employer_raw),
        employer=employer,
        employer_key=search_key(employer),
        address_raw=raw_text(address_raw),
        address=address,
        postal_code=postal_code(address),
        occupation_raw=raw_text(occupation_raw),
        noc_version="2021",
        noc_code=noc_code,
        occupation_title=occupation_title,
        occupation_label=f"{noc_code} — {occupation_title}",
        incorporate_status_raw=raw_text(incorporate_raw),
        incorporate_status=incorporate_status,
        lmia_measure=lmia_measure,
        lmia_count=integer(lmias_raw, "LMIA count", row_number),
        position_measure=position_measure,
        # The positive workbook contains one source row with zero positions.
        # Keep that published value verbatim and surface it as a quality note.
        position_count=integer(positions_raw, "position count", row_number, minimum=0),
        source_dataset_id=resource["dataset_id"],
        source_resource_id=resource["resource_id"],
        source_resource_url=resource["resource_page_url"],
        source_download_url=resource["download_url"],
        source_file=resource["filename"],
        source_sheet=resource["sheet"],
        source_row=row_number,
        source_sha256=resource["sha256"],
    )


def read_release(resource: dict[str, Any], raw_dir: Path) -> list[CanonicalRecord]:
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError as error:
        raise RuntimeError(
            "openpyxl 3.1.5 is required; install this project before running the pilot"
        ) from error

    workbook_path = raw_dir / resource["filename"]
    if not workbook_path.exists():
        raise FileNotFoundError(f"Official workbook not found: {workbook_path}")

    actual_hash = file_sha256(workbook_path)
    if actual_hash != resource["sha256"]:
        raise ValueError(
            f"Checksum mismatch for {resource['filename']}: "
            f"expected {resource['sha256']}, got {actual_hash}"
        )

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if workbook.sheetnames != [resource["sheet"]]:
            raise ValueError(
                f"Unexpected sheets in {resource['filename']}: {workbook.sheetnames!r}"
            )
        sheet = workbook[resource["sheet"]]
        title = clean_text(sheet.cell(1, 1).value)
        if resource["status"].title() not in title or "January to March 2025" not in title:
            raise ValueError(f"Unexpected title row in {resource['filename']}: {title!r}")
        headers = [clean_text(cell.value) for cell in sheet[2]]
        if headers != resource["expected_headers"]:
            raise ValueError(
                f"Header mismatch in {resource['filename']}: "
                f"expected {resource['expected_headers']!r}, got {headers!r}"
            )

        last_data_row = resource["expected_rows"] + 2
        note_label = clean_text(sheet.cell(last_data_row + 1, 1).value)
        if note_label != "Notes:":
            raise ValueError(
                f"Expected Notes: after row {last_data_row} in "
                f"{resource['filename']}, got {note_label!r}"
            )
        records = [
            normalize_row(resource, row_number, values)
            for row_number, values in enumerate(
                sheet.iter_rows(
                    min_row=3,
                    max_row=last_data_row,
                    max_col=8,
                    values_only=True,
                ),
                start=3,
            )
        ]
    finally:
        workbook.close()

    if len(records) != resource["expected_rows"]:
        raise ValueError(
            f"Row-count mismatch for {resource['filename']}: "
            f"expected {resource['expected_rows']}, got {len(records)}"
        )
    return records


def counter_rows(records: list[CanonicalRecord], key_name: str) -> list[dict[str, Any]]:
    buckets: dict[str, dict[str, Any]] = defaultdict(
        lambda: {
            "records": 0,
            "positive_records": 0,
            "negative_records": 0,
            "approved_lmias": 0,
            "requested_lmias": 0,
            "approved_positions": 0,
            "requested_positions_on_negative_decisions": 0,
        }
    )
    for record in records:
        key = getattr(record, key_name)
        bucket = buckets[key]
        bucket["records"] += 1
        bucket[f"{record.decision_status}_records"] += 1
        if record.decision_status == "positive":
            bucket["approved_lmias"] += record.lmia_count
            bucket["approved_positions"] += record.position_count
        else:
            bucket["requested_lmias"] += record.lmia_count
            bucket["requested_positions_on_negative_decisions"] += record.position_count
    return [
        {"name": name, **metrics}
        for name, metrics in sorted(
            buckets.items(), key=lambda item: (-item[1]["records"], item[0])
        )
    ]


def employer_rows(records: list[CanonicalRecord]) -> list[dict[str, Any]]:
    buckets: dict[str, dict[str, Any]] = {}
    for record in records:
        bucket = buckets.setdefault(
            record.employer_key,
            {
                "id": f"employer-{hashlib.sha256(record.employer_key.encode()).hexdigest()[:16]}",
                "search_key": record.employer_key,
                "variants": Counter(),
                "provinces": set(),
                "streams": set(),
                "occupations": set(),
                "positive_records": 0,
                "negative_records": 0,
                "approved_lmias": 0,
                "requested_lmias": 0,
                "approved_positions": 0,
                "requested_positions_on_negative_decisions": 0,
            },
        )
        bucket["variants"][record.employer] += 1
        bucket["provinces"].add(record.province_code)
        bucket["streams"].add(record.program_stream)
        bucket["occupations"].add(record.noc_code)
        bucket[f"{record.decision_status}_records"] += 1
        if record.decision_status == "positive":
            bucket["approved_lmias"] += record.lmia_count
            bucket["approved_positions"] += record.position_count
        else:
            bucket["requested_lmias"] += record.lmia_count
            bucket["requested_positions_on_negative_decisions"] += record.position_count

    result = []
    for bucket in buckets.values():
        variants = sorted(bucket.pop("variants").items(), key=lambda item: (-item[1], item[0]))
        result.append(
            {
                **bucket,
                "name": variants[0][0],
                "name_variants": [name for name, _ in variants],
                "provinces": sorted(bucket["provinces"]),
                "streams": sorted(bucket["streams"]),
                "occupations": sorted(bucket["occupations"]),
            }
        )
    return sorted(result, key=lambda row: row["name"].casefold())


def build_summary(records: list[CanonicalRecord], manifest: dict[str, Any]) -> dict[str, Any]:
    status_records = Counter(record.decision_status for record in records)
    approved = [record for record in records if record.decision_status == "positive"]
    negative = [record for record in records if record.decision_status == "negative"]
    employers = employer_rows(records)
    occupations = counter_rows(records, "occupation_label")
    provinces = counter_rows(records, "province_code")
    streams = counter_rows(records, "program_stream")

    return {
        "schema_version": "0.1.0",
        "source_data_as_of": max(
            resource.get("last_modified", resource["published_at"])
            for resource in manifest["resources"]
        ),
        "period": manifest["period"],
        "publication_status": "pilot",
        "totals": {
            "records": len(records),
            "positive_records": status_records["positive"],
            "negative_records": status_records["negative"],
            "employer_keys": len(employers),
            "approved_lmias": sum(record.lmia_count for record in approved),
            "requested_lmias": sum(record.lmia_count for record in negative),
            "approved_positions": sum(record.position_count for record in approved),
            "requested_positions_on_negative_decisions": sum(
                record.position_count for record in negative
            ),
        },
        "sources": [
            {
                "id": resource["id"],
                "status": resource["status"],
                "resource_page_url": resource["resource_page_url"],
                "download_url": resource["download_url"],
                "filename": resource["filename"],
                "sheet": resource["sheet"],
                "sha256": resource["sha256"],
                "rows": resource["expected_rows"],
                "published_at": resource["published_at"],
            }
            for resource in manifest["resources"]
        ],
        "limitations": manifest["limitations"],
        "quality_notes": [
            {
                "code": "zero-approved-positions",
                "records": sum(
                    record.decision_status == "positive" and record.position_count == 0
                    for record in records
                ),
                "text": "The official positive workbook contains one record with one approved LMIA and zero approved positions. The published zero is preserved.",
            }
        ],
        "provinces": provinces,
        "streams": streams,
        "top_occupations_by_approved_positions": sorted(
            occupations, key=lambda row: (-row["approved_positions"], row["name"])
        )[:15],
        "top_occupations_by_negative_positions": sorted(
            occupations,
            key=lambda row: (
                -row["requested_positions_on_negative_decisions"],
                row["name"],
            ),
        )[:15],
        "top_employers_by_approved_positions": sorted(
            employers, key=lambda row: (-row["approved_positions"], row["name"])
        )[:15],
        "top_employers_by_negative_positions": sorted(
            employers,
            key=lambda row: (
                -row["requested_positions_on_negative_decisions"],
                row["name"],
            ),
        )[:15],
        "checks": [
            "Both workbook SHA-256 checksums matched the official resource manifest.",
            "Workbook names, sheet names, titles, headers, and row counts matched expectations.",
            "Every canonical record has a stable ID and source workbook row.",
            "Positive and negative measures remain separate; no approval rate was inferred.",
        ],
    }


def write_outputs(
    records: list[CanonicalRecord], summary: dict[str, Any], output_root: Path
) -> None:
    normalized_dir = output_root / "data" / "normalized"
    published_dir = output_root / "data" / "published" / "2025q1"
    public_dir = output_root / "public" / "data" / "2025q1"
    normalized_dir.mkdir(parents=True, exist_ok=True)
    published_dir.mkdir(parents=True, exist_ok=True)
    public_dir.mkdir(parents=True, exist_ok=True)

    csv_path = normalized_dir / "lmia-2025q1.csv"
    field_names = [field.name for field in fields(CanonicalRecord)]
    with csv_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=field_names)
        writer.writeheader()
        writer.writerows(asdict(record) for record in records)

    employers = employer_rows(records)
    search_fields = [
        "id",
        "name",
        "search_key",
        "provinces",
        "positive_records",
        "negative_records",
        "approved_positions",
        "requested_positions_on_negative_decisions",
    ]
    search_rows = [[employer[field] for field in search_fields] for employer in employers]
    (published_dir / "summary.json").write_text(
        json.dumps(summary, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    (public_dir / "employers.json").write_text(
        json.dumps(
            {
                "schema_version": "0.1.0",
                "period": "2025-Q1",
                "fields": search_fields,
                "employers": search_rows,
            },
            ensure_ascii=False,
            separators=(",", ":"),
        )
        + "\n",
        encoding="utf-8",
    )


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("raw_dir", type=Path, help="Directory containing the official workbooks")
    parser.add_argument(
        "--output-root",
        type=Path,
        default=WORKSPACE_ROOT,
        help="Workspace receiving canonical and published outputs",
    )
    args = parser.parse_args()

    manifest = load_manifest()
    records: list[CanonicalRecord] = []
    for resource in manifest["resources"]:
        records.extend(read_release(resource, args.raw_dir.resolve()))
    summary = build_summary(records, manifest)
    write_outputs(records, summary, args.output_root.resolve())
    totals = summary["totals"]
    print(
        f"Normalized {totals['records']:,} records from {len(manifest['resources'])} "
        f"official resources; {totals['employer_keys']:,} normalized employer keys."
    )
    print(
        f"Positions: {totals['approved_positions']:,} approved; "
        f"{totals['requested_positions_on_negative_decisions']:,} requested on negative decisions."
    )


if __name__ == "__main__":
    main()
