"""Census downloaded LMIA employer-list CSV/XLS/XLSX resources.

Every workbook sheet is inspected. Header candidates and schema signatures are
reported rather than silently selecting the first sheet or coercing an unknown
layout into the current-quarter schema.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import re
import shutil
import subprocess
import tempfile
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Iterable, Iterator

from lmiacsv.scrape_catalogue import CatalogueError, write_json


WORKSPACE_ROOT = Path(__file__).resolve().parents[1]
MANIFEST_PATH = WORKSPACE_ROOT / "data" / "raw" / "employer-lists" / "manifest.json"
DEFAULT_OUTPUT_PATH = (
    WORKSPACE_ROOT / "data" / "published" / "catalogue" / "employer-resource-census.json"
)
MAX_SCAN_ROWS = 250_000
MAX_SCAN_COLUMNS = 50
HEADER_SCAN_ROWS = 30
HEADER_TERMS = (
    "employer",
    "province",
    "territory",
    "program stream",
    "stream",
    "occupation",
    "noc",
    "address",
    "location",
    "positions",
    "labour market impact assessment",
    "lmia",
    "incorporate",
)
NOTE_RE = re.compile(r"^(notes?|remarques?)\s*:?$", re.I)


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(unicodedata.normalize("NFKC", str(value)).strip().split())


def normalized_header(value: Any) -> str:
    text = unicodedata.normalize("NFKD", clean_text(value).casefold())
    text = "".join(character for character in text if not unicodedata.combining(character))
    return " ".join(re.sub(r"[^a-z0-9]+", " ", text).split())


def trim_row(values: Iterable[Any]) -> list[str]:
    row = [clean_text(value) for value in list(values)[:MAX_SCAN_COLUMNS]]
    while row and not row[-1]:
        row.pop()
    return row


def header_candidate(rows: list[list[str]]) -> dict[str, Any] | None:
    candidates = []
    for row_number, row in enumerate(rows[:HEADER_SCAN_ROWS], start=1):
        normalized = [normalized_header(value) for value in row]
        nonempty = [value for value in normalized if value]
        matched_terms = sorted(
            {term for term in HEADER_TERMS if any(term in value for value in nonempty)}
        )
        if len(matched_terms) < 2 or len(nonempty) < 2:
            continue
        score = len(matched_terms) * 10 + min(len(nonempty), 10)
        candidates.append(
            {
                "row": row_number,
                "score": score,
                "matched_terms": matched_terms,
                "values": row,
                "normalized_values": normalized,
            }
        )
    if not candidates:
        return None
    return max(candidates, key=lambda candidate: (candidate["score"], -candidate["row"]))


def analyze_rows(rows: Iterator[Iterable[Any]]) -> dict[str, Any]:
    leading: list[list[str]] = []
    nonempty_rows: list[int] = []
    note_rows: list[int] = []
    row_widths: list[int] = []
    scanned_rows = 0
    truncated = False
    for row_number, values in enumerate(rows, start=1):
        if row_number > MAX_SCAN_ROWS:
            truncated = True
            break
        scanned_rows = row_number
        row = trim_row(values)
        if row_number <= HEADER_SCAN_ROWS:
            leading.append(row)
        if any(row):
            nonempty_rows.append(row_number)
            row_widths.append(sum(bool(value) for value in row))
            if row and NOTE_RE.match(row[0]):
                note_rows.append(row_number)

    header = header_candidate(leading)
    first_note_after_header = None
    estimated_data_rows = None
    if header:
        first_note_after_header = next((row for row in note_rows if row > header["row"]), None)
        data_end = first_note_after_header or scanned_rows + 1
        estimated_data_rows = sum(
            header["row"] < row_number < data_end for row_number in nonempty_rows
        )
    title_rows = []
    if header:
        title_rows = [
            {"row": index, "values": row}
            for index, row in enumerate(leading[: header["row"] - 1], start=1)
            if any(row)
        ]
    return {
        "scanned_rows": scanned_rows,
        "scan_truncated": truncated,
        "nonempty_rows": len(nonempty_rows),
        "maximum_nonempty_cells": max(row_widths, default=0),
        "header": header,
        "title_rows": title_rows,
        "first_note_row": first_note_after_header,
        "estimated_data_rows": estimated_data_rows,
    }


def decode_csv(path: Path) -> tuple[str, str]:
    content = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return content.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    raise CatalogueError(f"Unable to decode CSV: {path}")


def inspect_csv(path: Path) -> dict[str, Any]:
    text, encoding = decode_csv(path)
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","
    analysis = analyze_rows(iter(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)))
    return {
        "kind": "csv",
        "encoding": encoding,
        "delimiter": delimiter,
        "sheet_count": 1,
        "sheets": [
            {
                "name": "CSV",
                "state": "visible",
                "max_row": analysis["scanned_rows"],
                "max_column": analysis["maximum_nonempty_cells"],
                "merged_range_count": 0,
                **analysis,
            }
        ],
    }


def soffice_path() -> str:
    executable = shutil.which("soffice") or shutil.which("libreoffice")
    if not executable:
        raise RuntimeError("LibreOffice soffice is required to inspect legacy XLS files")
    return executable


def convert_xls(path: Path, output_dir: Path) -> Path:
    command = [
        soffice_path(),
        "--headless",
        "--convert-to",
        "xlsx",
        "--outdir",
        str(output_dir),
        str(path),
    ]
    result = subprocess.run(command, capture_output=True, text=True, timeout=120)
    converted = output_dir / f"{path.stem}.xlsx"
    if result.returncode != 0 or not converted.exists():
        raise RuntimeError(
            f"LibreOffice failed to convert {path.name}: "
            f"{result.stderr.strip() or result.stdout.strip()}"
        )
    return converted


def inspect_xlsx(path: Path, converted_from_xls: bool = False) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError as error:
        raise RuntimeError("openpyxl 3.1.5 is required for the workbook census") from error

    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        sheets = []
        for sheet in workbook.worksheets:
            max_row = min(sheet.max_row or 0, MAX_SCAN_ROWS)
            max_column = min(sheet.max_column or 0, MAX_SCAN_COLUMNS)
            rows = sheet.iter_rows(
                min_row=1,
                max_row=max_row,
                max_col=max_column,
                values_only=True,
            )
            analysis = analyze_rows(iter(rows))
            sheets.append(
                {
                    "name": sheet.title,
                    "state": sheet.sheet_state,
                    "max_row": sheet.max_row,
                    "max_column": sheet.max_column,
                    "merged_range_count": len(sheet.merged_cells.ranges),
                    "hidden_row_count": sum(
                        dimension.hidden is True for dimension in sheet.row_dimensions.values()
                    ),
                    "hidden_column_count": sum(
                        dimension.hidden is True for dimension in sheet.column_dimensions.values()
                    ),
                    **analysis,
                }
            )
    finally:
        workbook.close()
    return {
        "kind": "xls_converted_for_inspection" if converted_from_xls else "xlsx",
        "sheet_count": len(sheets),
        "sheets": sheets,
    }


def inspect_resource(path: Path, format_name: str) -> dict[str, Any]:
    if format_name == "CSV":
        return inspect_csv(path)
    if format_name == "XLSX":
        return inspect_xlsx(path)
    if format_name == "XLS":
        with tempfile.TemporaryDirectory(prefix="lmia-xls-census-") as directory:
            converted = convert_xls(path, Path(directory))
            return inspect_xlsx(converted, converted_from_xls=True)
    raise CatalogueError(f"Unsupported census format: {format_name}")


def schema_signature(header: dict[str, Any]) -> str:
    values = header["normalized_values"]
    while values and not values[-1]:
        values = values[:-1]
    content = json.dumps(values, ensure_ascii=False, separators=(",", ":")).encode()
    return hashlib.sha256(content).hexdigest()[:16]


def decision_signal(header: dict[str, Any]) -> str | None:
    values = header["normalized_values"]
    joined = " | ".join(values)
    if (
        "approved lmia" in joined
        or "approved positions" in joined
        or "positions approved" in joined
    ):
        return "positive"
    if (
        "requested lmia" in joined
        or "requested positions" in joined
        or "positions requested" in joined
    ):
        return "negative"
    return None


def choose_sheet(sheets: list[dict[str, Any]]) -> dict[str, Any] | None:
    candidates = [sheet for sheet in sheets if sheet.get("header")]
    if not candidates:
        return None
    return max(
        candidates,
        key=lambda sheet: (
            sheet["state"] == "visible",
            sheet["header"]["score"],
            sheet.get("estimated_data_rows") or 0,
            -sheet["header"]["row"],
        ),
    )


def resolve_archive_path(value: str) -> Path:
    path = Path(value)
    return path if path.is_absolute() else WORKSPACE_ROOT / path


def build_census(manifest: dict[str, Any]) -> dict[str, Any]:
    resources = []
    failures = []
    for entry in manifest.get("resources", []):
        if entry.get("status") != "downloaded":
            failures.append(
                {
                    "resource_id": entry.get("resource_id"),
                    "stage": "download",
                    "error": entry.get("error", "Resource was not downloaded"),
                }
            )
            continue
        path = resolve_archive_path(entry["archive_path"])
        try:
            inspection = inspect_resource(path, entry["format"])
        except Exception as error:
            failures.append(
                {
                    "resource_id": entry["resource_id"],
                    "stage": "inspection",
                    "error": f"{type(error).__name__}: {error}",
                }
            )
            continue
        selected = choose_sheet(inspection["sheets"])
        candidate_sheets = [sheet for sheet in inspection["sheets"] if sheet.get("header")]
        resource = {
            **{
                key: entry[key]
                for key in (
                    "dataset_id",
                    "decision",
                    "resource_id",
                    "language",
                    "period",
                    "format",
                    "catalogue_format",
                    "format_matches_catalogue",
                    "source_filename",
                    "catalogue_modified_at",
                    "downloaded_bytes",
                    "size_matches_catalogue",
                    "sha256",
                    "archive_path",
                )
            },
            "inspection_kind": inspection["kind"],
            "text_encoding": inspection.get("encoding"),
            "delimiter": inspection.get("delimiter"),
            "sheet_count": inspection["sheet_count"],
            "candidate_sheet_count": len(candidate_sheets),
            "selected_sheet": selected["name"] if selected else None,
            "selected_schema_signature": schema_signature(selected["header"]) if selected else None,
            "header_decision_signal": decision_signal(selected["header"]) if selected else None,
            "sheets": inspection["sheets"],
        }
        resources.append(resource)

    clusters: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for resource in resources:
        signature = resource.get("selected_schema_signature")
        if signature:
            clusters[signature].append(resource)
    schema_clusters = []
    for signature, members in sorted(clusters.items()):
        example = members[0]
        selected = next(
            sheet for sheet in example["sheets"] if sheet["name"] == example["selected_sheet"]
        )
        schema_clusters.append(
            {
                "signature": signature,
                "resource_count": len(members),
                "decisions": dict(
                    sorted(Counter(member["decision"] for member in members).items())
                ),
                "formats": dict(sorted(Counter(member["format"] for member in members).items())),
                "periods": sorted({member["period"]["key"] for member in members}),
                "header_values": selected["header"]["values"],
                "normalized_header_values": selected["header"]["normalized_values"],
                "resource_ids": sorted(member["resource_id"] for member in members),
            }
        )
    schema_clusters.sort(key=lambda cluster: (-cluster["resource_count"], cluster["signature"]))

    content_groups: dict[str, list[str]] = defaultdict(list)
    for resource in resources:
        content_groups[resource["sha256"]].append(resource["resource_id"])
    resources_by_id = {resource["resource_id"]: resource for resource in resources}
    duplicate_content = [
        {
            "sha256": digest,
            "resource_ids": sorted(resource_ids),
            "decisions": sorted(
                {resources_by_id[resource_id]["decision"] for resource_id in resource_ids}
            ),
            "periods": sorted(
                {resources_by_id[resource_id]["period"]["key"] for resource_id in resource_ids}
            ),
            "source_filenames": sorted(
                {resources_by_id[resource_id]["source_filename"] for resource_id in resource_ids}
            ),
        }
        for digest, resource_ids in sorted(content_groups.items())
        if len(resource_ids) > 1
    ]
    return {
        "schema_version": "0.1.0",
        # Retrieval time remains in the ignored raw manifest. The tracked census
        # fingerprints content and source metadata, so unchanged inputs remain stable.
        "source_manifest_selection": manifest.get("selection"),
        "resource_count": len(manifest.get("resources", [])),
        "inspected_count": len(resources),
        "failure_count": len(failures),
        "schema_cluster_count": len(schema_clusters),
        "counts": {
            "formats": dict(sorted(Counter(resource["format"] for resource in resources).items())),
            "decisions": dict(
                sorted(Counter(resource["decision"] for resource in resources).items())
            ),
            "inspection_kinds": dict(
                sorted(Counter(resource["inspection_kind"] for resource in resources).items())
            ),
            "text_encodings": dict(
                sorted(
                    Counter(
                        resource["text_encoding"]
                        for resource in resources
                        if resource["text_encoding"]
                    ).items()
                )
            ),
        },
        "resources": resources,
        "schema_clusters": schema_clusters,
        "audit": {
            "failures": failures,
            "resources_without_header_candidate": sorted(
                resource["resource_id"] for resource in resources if not resource["selected_sheet"]
            ),
            "resources_with_multiple_candidate_sheets": sorted(
                resource["resource_id"]
                for resource in resources
                if resource["candidate_sheet_count"] > 1
            ),
            "resources_with_hidden_sheets": sorted(
                resource["resource_id"]
                for resource in resources
                if any(sheet["state"] != "visible" for sheet in resource["sheets"])
            ),
            "resources_with_size_mismatch": sorted(
                resource["resource_id"]
                for resource in resources
                if not resource["size_matches_catalogue"]
            ),
            "resources_with_format_mismatch": sorted(
                resource["resource_id"]
                for resource in resources
                if not resource["format_matches_catalogue"]
            ),
            "resources_with_decision_header_mismatch": sorted(
                resource["resource_id"]
                for resource in resources
                if resource["header_decision_signal"]
                and resource["header_decision_signal"] != resource["decision"]
            ),
            "resources_with_latin1_fallback": sorted(
                resource["resource_id"]
                for resource in resources
                if resource["text_encoding"] == "latin-1"
            ),
            "duplicate_content": duplicate_content,
        },
        "checks": [
            "Every downloaded blob is addressed by dataset ID, resource UUID and SHA-256.",
            "Every workbook sheet is inspected; the first sheet is never selected by position alone.",
            "Header candidates require multiple LMIA-domain terms and retain their original values.",
            "Schema clusters use exact normalized header signatures; unknown shapes remain visible.",
            "Legacy XLS files are converted only in a temporary directory for inspection; originals remain immutable.",
        ],
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--manifest", type=Path, default=MANIFEST_PATH)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT_PATH)
    args = parser.parse_args()
    manifest = json.loads(args.manifest.read_text(encoding="utf-8"))
    census = build_census(manifest)
    write_json(args.output.resolve(), census)
    print(
        f"Inspected {census['inspected_count']}/{census['resource_count']} resources; "
        f"found {census['schema_cluster_count']} exact header-signature clusters."
    )
    print(
        f"Hidden sheets: {len(census['audit']['resources_with_hidden_sheets'])}; "
        f"multiple candidates: {len(census['audit']['resources_with_multiple_candidate_sheets'])}; "
        f"failures: {census['failure_count']}."
    )
    if census["failure_count"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
