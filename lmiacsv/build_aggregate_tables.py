"""Normalize archived aggregate Table 14 sources into auditable CSV products.

Every source workbook is parsed independently. Overlapping quarterly releases are
compared cell by cell; the newest release wins only after revisions are recorded.
Complete quarterly years are reconciled to the separate official annual Table 14.
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import re
import unicodedata
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from lmiacsv.fetch_aggregate_resources import FIRST_EXPECTED_PERIOD, period_range, period_tuple
from lmiacsv.scrape_catalogue import CatalogueError, atomic_write, write_json


WORKSPACE_ROOT = Path(__file__).resolve().parents[1]
MANIFEST_PATH = WORKSPACE_ROOT / "data" / "raw" / "aggregate-tables" / "manifest.json"
OUTPUT_ROOT = WORKSPACE_ROOT / "data" / "full"
QUARTERLY_OUTPUT = OUTPUT_ROOT / "positive-quarterly.csv"
ANNUAL_OUTPUT = OUTPUT_ROOT / "positive-annual.csv"
NATIONAL_OUTPUT = OUTPUT_ROOT / "positive-quarterly-national.csv"
META_ROOT = WORKSPACE_ROOT / "data" / "meta"
SCHEMA_OUTPUT = META_ROOT / "aggregates.schema.json"
RECONCILIATION_OUTPUT = META_ROOT / "aggregates.reconciliation.json"

PROVINCES = {
    "newfoundland and labrador": "NL",
    "newfoundland labrador": "NL",
    "terre neuve et labrador": "NL",
    "nl": "NL",
    "prince edward island": "PE",
    "ile du prince edouard": "PE",
    "le du prince douard": "PE",
    "pe": "PE",
    "nova scotia": "NS",
    "nouvelle ecosse": "NS",
    "nouvelle cosse": "NS",
    "ns": "NS",
    "new brunswick": "NB",
    "nouveau brunswick": "NB",
    "nb": "NB",
    "quebec": "QC",
    "québec": "QC",
    "qc": "QC",
    "ontario": "ON",
    "on": "ON",
    "manitoba": "MB",
    "mb": "MB",
    "saskatchewan": "SK",
    "sk": "SK",
    "alberta": "AB",
    "ab": "AB",
    "british columbia": "BC",
    "colombie britannique": "BC",
    "bc": "BC",
    "yukon": "YT",
    "yt": "YT",
    "northwest territories": "NT",
    "northwest territory": "NT",
    "territoires du nord ouest": "NT",
    "nt": "NT",
    "nunavut": "NU",
    "nu": "NU",
}
STREAMS = {
    "high wage": "high_wage",
    "haut salaire": "high_wage",
    "low wage": "low_wage",
    "bas salaire": "low_wage",
    "primary agriculture": "primary_agriculture",
    "agriculture primaire": "primary_agriculture",
    "global talent stream": "global_talent_stream",
    "volet des talents mondiaux": "global_talent_stream",
    "volet des talents mondial": "global_talent_stream",
    "talent mondial": "global_talent_stream",
    "permanent resident only": "permanent_resident_only",
    "permanent residence only": "permanent_resident_only",
    "permanent residency only": "permanent_resident_only",
    "residence permanente uniquement": "permanent_resident_only",
    "resident permanent seulement": "permanent_resident_only",
    "caregiver": "caregiver",
    "caregivers": "caregiver",
    "aidant familial": "caregiver",
    "aidants familiaux": "caregiver",
    "unspecified": "unspecified",
    "non precise": "unspecified",
    "non specifie": "unspecified",
}
MONTH_RANGE_ALIASES = (
    ({"jan", "janv", "janvier"}, {"mar", "mars"}),
    ({"apr", "avr", "avril"}, {"jun", "juin", "jui"}),
    ({"jul", "juil", "juillet", "jui"}, {"sep", "sept", "septembre"}),
    ({"oct", "octobre"}, {"dec", "decembre"}),
)


class AggregateTableError(CatalogueError):
    """Raised when an aggregate workbook does not meet the declared contract."""


def clean_text(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def search_text(value: Any) -> str:
    folded = unicodedata.normalize("NFKD", clean_text(value).casefold())
    characters = (
        character if character.isalnum() else " "
        for character in folded
        if not unicodedata.combining(character)
    )
    return " ".join("".join(characters).split())


def header_text(value: Any) -> str:
    return clean_text(value).casefold().replace(" / ", "/")


def normalize_province(value: Any) -> tuple[str, str]:
    raw = clean_text(value)
    key = search_text(raw)
    if key not in PROVINCES:
        raise AggregateTableError(f"Unknown province/territory label {raw!r}")
    return PROVINCES[key], raw


def normalize_stream(value: Any) -> tuple[str, str]:
    raw = clean_text(value)
    key = search_text(raw)
    if key not in STREAMS:
        raise AggregateTableError(f"Unknown program stream label {raw!r}")
    return STREAMS[key], raw


def quarterly_period(value: Any) -> str | None:
    text = search_text(value)
    direct = re.fullmatch(r"(\d{4})\s*[qt]\s*([1-4])", text, re.I)
    if direct:
        return f"{direct.group(1)}Q{direct.group(2)}"
    reverse = re.fullmatch(r"[qt]\s*([1-4])\s*(\d{4})", text, re.I)
    if reverse:
        return f"{reverse.group(2)}Q{reverse.group(1)}"
    months = re.fullmatch(r"([a-z]+)\s+([a-z]+)\s*(\d{4})", text, re.I)
    if months:
        quarter = month_range_quarter(months.group(1), months.group(2))
        if quarter:
            return f"{months.group(3)}Q{quarter}"
    return None


def month_range_quarter(start: str, end: str) -> int | None:
    for quarter, (start_aliases, end_aliases) in enumerate(MONTH_RANGE_ALIASES, 1):
        if start.lower() in start_aliases and end.lower() in end_aliases:
            return quarter
    return None


def quarter_from_month_range(value: Any) -> int | None:
    text = search_text(value)
    months = re.fullmatch(r"([a-z]+)\s+([a-z]+)", text, re.I)
    if not months:
        return None
    return month_range_quarter(months.group(1), months.group(2))


def annual_period(value: Any) -> int | None:
    if isinstance(value, int) and 1900 <= value <= 2200:
        return value
    if isinstance(value, float) and value.is_integer() and 1900 <= value <= 2200:
        return int(value)
    text = clean_text(value)
    return int(text) if re.fullmatch(r"20\d{2}", text) else None


def integer_count(value: Any, context: str) -> int:
    if isinstance(value, bool):
        raise AggregateTableError(f"Boolean count in {context}")
    if isinstance(value, int):
        result = value
    elif isinstance(value, float) and value.is_integer():
        result = int(value)
    elif isinstance(value, str) and re.fullmatch(r"\d[\d,\s\u00a0\u202f]*", value.strip()):
        result = int(re.sub(r"[,\s\u00a0\u202f]", "", value))
    else:
        raise AggregateTableError(f"Expected an integer count in {context}, found {value!r}")
    if result < 0:
        raise AggregateTableError(f"Negative LMIA count in {context}: {result}")
    return result


def read_csv_matrix(path: Path) -> tuple[str, list[list[Any]]]:
    content = path.read_bytes()
    text = None
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise AggregateTableError(f"Could not decode CSV source {path}")
    try:
        dialect = csv.Sniffer().sniff(text[:65536], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return "CSV", [list(row) for row in csv.reader(io.StringIO(text), dialect)]


def read_xlsx_matrices(path: Path) -> list[tuple[str, list[list[Any]]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    values = []
    for sheet in workbook.worksheets:
        if sheet.sheet_state != "visible":
            continue
        matrix = [list(row) for row in sheet.iter_rows(values_only=True)]
        if any(any(clean_text(cell) for cell in row) for row in matrix):
            values.append((sheet.title, matrix))
    workbook.close()
    return values


def source_matrices(path: Path, format_name: str) -> list[tuple[str, list[list[Any]]]]:
    if format_name == "CSV":
        return [read_csv_matrix(path)]
    if format_name == "XLSX":
        return read_xlsx_matrices(path)
    raise AggregateTableError(f"Unsupported aggregate source format {format_name!r}")


def header_row(matrix: list[list[Any]]) -> tuple[int, list[Any]]:
    for index, row in enumerate(matrix[:30]):
        headers = [header_text(value) for value in row]
        has_province = any(
            value.startswith("province")
            and "territo" in value
            for value in headers
        )
        has_stream = any(
            value in {"program stream", "stream", "volet", "volet du programme"}
            for value in headers
        )
        if has_province and has_stream:
            return index, row
    raise AggregateTableError("No Table 14 province/territory and program-stream header found")


def parse_matrix(
    matrix: list[list[Any]], sheet_name: str, source: dict[str, Any]
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    header_index, headers = header_row(matrix)
    normalized_headers = [header_text(value) for value in headers]
    province_indexes = [
        index
        for index, value in enumerate(normalized_headers)
        if "province" in value and "territo" in value
    ]
    stream_indexes = [
        index
        for index, value in enumerate(normalized_headers)
        if "stream" in value or "volet" in value
    ]
    if len(province_indexes) != 1 or len(stream_indexes) != 1:
        raise AggregateTableError("Table 14 must have one province and one stream column")
    province_index = province_indexes[0]
    stream_index = stream_indexes[0]

    period_columns: list[tuple[int, str | int]] = []
    inherited_years: dict[int, int] = {}
    if source["source_kind"] == "quarterly_release" and header_index:
        for prior_row in matrix[max(0, header_index - 3) : header_index]:
            active_year = None
            for index in range(len(headers)):
                value = prior_row[index] if index < len(prior_row) else None
                explicit_year = annual_period(value)
                if explicit_year is not None:
                    active_year = explicit_year
                if active_year is not None:
                    inherited_years[index] = active_year
    for index, value in enumerate(headers):
        period = (
            quarterly_period(value)
            if source["source_kind"] == "quarterly_release"
            else annual_period(value)
        )
        if period is None and source["source_kind"] == "quarterly_release":
            quarter = quarter_from_month_range(value)
            year = inherited_years.get(index)
            if quarter is not None and year is not None:
                period = f"{year}Q{quarter}"
        if period is not None:
            period_columns.append((index, period))
    if not period_columns:
        raise AggregateTableError(f"No period columns found in {source['source_filename']}")
    if len({period for _index, period in period_columns}) != len(period_columns):
        raise AggregateTableError(f"Duplicate period columns in {source['source_filename']}")

    observations = []
    ignored_note_rows = []
    ignored_summary_rows = []
    seen_dimensions = set()
    for row_index, row in enumerate(matrix[header_index + 1 :], start=header_index + 2):
        province_value = row[province_index] if province_index < len(row) else None
        stream_value = row[stream_index] if stream_index < len(row) else None
        if not clean_text(province_value) and not clean_text(stream_value):
            if any(clean_text(value) for value in row):
                ignored_note_rows.append(row_index)
            continue
        if not clean_text(province_value) or not clean_text(stream_value):
            ignored_note_rows.append(row_index)
            continue
        if search_text(province_value) == "canada" or search_text(stream_value) in {
            "total",
            "all streams",
            "tous les volets",
        }:
            ignored_summary_rows.append(row_index)
            continue
        province, province_raw = normalize_province(province_value)
        stream, stream_raw = normalize_stream(stream_value)
        dimension = (province, stream)
        if dimension in seen_dimensions:
            raise AggregateTableError(
                f"Duplicate province/stream row {dimension!r} in {source['source_filename']}"
            )
        seen_dimensions.add(dimension)
        for column_index, period in period_columns:
            value = row[column_index] if column_index < len(row) else None
            count = integer_count(
                value,
                f"{source['source_filename']} sheet {sheet_name!r} row {row_index} period {period}",
            )
            observation = {
                "province_territory": province,
                "province_territory_raw": province_raw,
                "program_stream": stream,
                "program_stream_raw": stream_raw,
                "positive_lmias": count,
                "source_dataset_id": source["dataset_id"],
                "source_resource_id": source["resource_id"],
                "source_language": source.get("language", "en"),
                "source_release_period": source.get("release_period"),
                "source_filename": source["source_filename"],
                "source_sha256": source["sha256"],
                "source_sheet": sheet_name,
                "source_row": row_index,
            }
            if source["source_kind"] == "quarterly_release":
                year, quarter = period_tuple(str(period))
                observation.update({"period": period, "year": year, "quarter": quarter})
            else:
                observation.update({"year": period})
            observations.append(observation)
    if not observations:
        raise AggregateTableError(f"No Table 14 observations in {source['source_filename']}")
    return observations, {
        "sheet": sheet_name,
        "header_row": header_index + 1,
        "data_dimension_count": len(seen_dimensions),
        "observation_count": len(observations),
        "periods": [str(period) for _index, period in period_columns],
        "ignored_note_rows": ignored_note_rows,
        "ignored_summary_rows": ignored_summary_rows,
    }


def parse_source(source: dict[str, Any]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    path = Path(source["archive_path"])
    if not path.is_absolute():
        path = WORKSPACE_ROOT / path
    if not path.is_file():
        raise AggregateTableError(f"Aggregate source is missing: {path}")
    parsed = []
    failures = []
    for sheet_name, matrix in source_matrices(path, source["format"]):
        try:
            parsed.append(parse_matrix(matrix, sheet_name, source))
        except AggregateTableError as error:
            failures.append({"sheet": sheet_name, "error": str(error)})
    if len(parsed) != 1:
        raise AggregateTableError(
            f"Expected one recognizable Table 14 sheet in {source['source_filename']}; "
            f"recognized={len(parsed)}, failures={failures!r}"
        )
    observations, details = parsed[0]
    return observations, {
        "dataset_id": source["dataset_id"],
        "resource_id": source["resource_id"],
        "language": source.get("language", "en"),
        "release_period": source.get("release_period"),
        "source_filename": source["source_filename"],
        "sha256": source["sha256"],
        "format": source["format"],
        **details,
    }


def latest_quarterly(
    observations: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], set[tuple[str, str, str]]]:
    grouped: dict[tuple[str, str, str], list[dict[str, Any]]] = defaultdict(list)
    for observation in observations:
        key = (
            observation["period"],
            observation["province_territory"],
            observation["program_stream"],
        )
        grouped[key].append(observation)
    selected = []
    revisions = []
    revised_keys = set()
    for key, values in grouped.items():
        values.sort(
            key=lambda value: (
                period_tuple(value["source_release_period"]),
                value["source_sha256"],
                value["source_filename"],
            )
        )
        chosen = values[-1]
        selected.append(chosen)
        distinct = {value["positive_lmias"] for value in values}
        if len(distinct) > 1:
            revised_keys.add(key)
            revisions.append(
                {
                    "period": key[0],
                    "province_territory": key[1],
                    "program_stream": key[2],
                    "selected_value": chosen["positive_lmias"],
                    "selected_release": chosen["source_release_period"],
                    "observations": [
                        {
                            "release_period": value["source_release_period"],
                            "positive_lmias": value["positive_lmias"],
                            "source_sha256": value["source_sha256"],
                        }
                        for value in values
                    ],
                }
            )
    selected.sort(
        key=lambda value: (
            period_tuple(value["period"]),
            value["province_territory"],
            value["program_stream"],
        )
    )
    revisions.sort(
        key=lambda value: (
            period_tuple(value["period"]),
            value["province_territory"],
            value["program_stream"],
        )
    )
    return selected, revisions, revised_keys


def annual_reconciliation(
    quarterly: list[dict[str, Any]], annual: list[dict[str, Any]]
) -> list[dict[str, Any]]:
    quarter_groups: dict[tuple[int, str, str], list[dict[str, Any]]] = defaultdict(list)
    for observation in quarterly:
        quarter_groups[
            (
                observation["year"],
                observation["province_territory"],
                observation["program_stream"],
            )
        ].append(observation)
    annual_values = {
        (value["year"], value["province_territory"], value["program_stream"]): value
        for value in annual
    }
    periods_by_year: dict[int, set[str]] = defaultdict(set)
    for observation in quarterly:
        periods_by_year[observation["year"]].add(observation["period"])
    complete_years = {
        year
        for year, periods in periods_by_year.items()
        if periods == {f"{year}Q{quarter}" for quarter in range(1, 5)}
    }
    results = []
    for key, quarters in sorted(quarter_groups.items()):
        periods = sorted({value["period"] for value in quarters}, key=period_tuple)
        if len(periods) != 4:
            continue
        annual_value = annual_values.get(key)
        quarterly_total = sum(value["positive_lmias"] for value in quarters)
        if annual_value is None:
            results.append(
                {
                    "year": key[0],
                    "province_territory": key[1],
                    "program_stream": key[2],
                    "status": "missing_annual_observation",
                    "quarterly_total": quarterly_total,
                    "annual_total": None,
                    "difference": None,
                }
            )
            continue
        annual_total = annual_value["positive_lmias"]
        results.append(
            {
                "year": key[0],
                "province_territory": key[1],
                "program_stream": key[2],
                "status": "match" if quarterly_total == annual_total else "mismatch",
                "quarterly_total": quarterly_total,
                "annual_total": annual_total,
                "difference": quarterly_total - annual_total,
            }
        )
    for key, annual_value in annual_values.items():
        if key[0] not in complete_years or key in quarter_groups:
            continue
        results.append(
            {
                "year": key[0],
                "province_territory": key[1],
                "program_stream": key[2],
                "status": "missing_quarterly_observation",
                "quarterly_total": None,
                "annual_total": annual_value["positive_lmias"],
                "difference": None,
            }
        )
    results.sort(
        key=lambda value: (
            value["year"],
            value["province_territory"],
            value["program_stream"],
        )
    )
    return results


def annual_national_summary(checks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[int, list[dict[str, Any]]] = defaultdict(list)
    for check in checks:
        grouped[check["year"]].append(check)
    values = []
    for year, rows in sorted(grouped.items()):
        quarterly_total = sum(
            row["quarterly_total"]
            for row in rows
            if row["quarterly_total"] is not None
        )
        annual_total = sum(
            row["annual_total"] for row in rows if row["annual_total"] is not None
        )
        annual_only = [
            row for row in rows if row["status"] == "missing_quarterly_observation"
        ]
        quarterly_only = [
            row for row in rows if row["status"] == "missing_annual_observation"
        ]
        values.append(
            {
                "year": year,
                "status": "match" if quarterly_total == annual_total else "mismatch",
                "quarterly_total": quarterly_total,
                "annual_total": annual_total,
                "difference": quarterly_total - annual_total,
                "common_scope_difference": sum(
                    row["difference"] for row in rows if row["difference"] is not None
                ),
                "annual_only_category_count": len(annual_only),
                "annual_only_positive_lmias": sum(row["annual_total"] for row in annual_only),
                "quarterly_only_category_count": len(quarterly_only),
                "quarterly_only_positive_lmias": sum(
                    row["quarterly_total"] for row in quarterly_only
                ),
            }
        )
    return values


def csv_bytes(rows: list[dict[str, Any]], fieldnames: list[str]) -> bytes:
    output = io.StringIO(newline="")
    writer = csv.DictWriter(output, fieldnames=fieldnames, lineterminator="\n")
    writer.writeheader()
    writer.writerows({field: row.get(field) for field in fieldnames} for row in rows)
    return output.getvalue().encode("utf-8")


def national_rows(
    quarterly: list[dict[str, Any]],
    revised_keys: set[tuple[str, str, str]],
    annual_checks: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in quarterly:
        grouped[row["period"]].append(row)
    reconciliation_flags: dict[int, set[str]] = defaultdict(set)
    for check in annual_checks:
        if check["status"] == "mismatch":
            reconciliation_flags[check["year"]].add("annual_reconciliation_mismatch")
        if (
            check["status"] == "missing_quarterly_observation"
            and check["annual_total"] > 0
        ):
            reconciliation_flags[check["year"]].add("annual_scope_has_unallocated_lmias")
    values = []
    for period in sorted(grouped, key=period_tuple):
        rows = grouped[period]
        year, quarter = period_tuple(period)
        releases = sorted({row["source_release_period"] for row in rows}, key=period_tuple)
        revised = any(
            (period, row["province_territory"], row["program_stream"]) in revised_keys
            for row in rows
        )
        flags = set(reconciliation_flags.get(year, set()))
        if revised:
            flags.add("revised_across_releases")
        values.append(
            {
                "period": period,
                "year": year,
                "quarter": quarter,
                "positive_lmias": sum(row["positive_lmias"] for row in rows),
                "contributing_rows": len(rows),
                "source_release_periods": ";".join(releases),
                "quality_flags": ";".join(sorted(flags)),
            }
        )
    return values


def schema_document() -> dict[str, Any]:
    return {
        "schema_version": "0.2.0",
        "dimensions": {
            "province_territory": "stable Canadian postal abbreviation or territory code",
            "program_stream": "stable language-neutral key joined to translations.csv",
            "source_language": "language of the parsed official source; normalized products use en",
        },
        "products": {
            "positive-quarterly.csv": {
                "grain": "decision quarter × province/territory × program stream",
                "measure": {
                    "column": "positive_lmias",
                    "unit": "Labour Market Impact Assessments",
                    "scope": "positive decisions",
                    "period_basis": "LMIA decision date",
                },
                "source": "Official ESDC quarterly Table 14",
            },
            "positive-annual.csv": {
                "grain": "decision year × province/territory × program stream",
                "measure": {
                    "column": "positive_lmias",
                    "unit": "Labour Market Impact Assessments",
                    "scope": "positive decisions",
                    "period_basis": "LMIA decision date",
                },
                "source": "Official ESDC annual Table 14",
            },
            "positive-quarterly-national.csv": {
                "grain": "decision quarter",
                "measure": {
                    "column": "positive_lmias",
                    "formula": "sum of normalized Table 14 province/stream observations",
                },
            },
        },
        "non_claims": [
            "Positive LMIAs are not approved positions, work permits, workers, or employers.",
            "This product does not contain complete all-stream negative LMIA decisions.",
            "No national LMIA approval or denial rate can be calculated from these files.",
        ],
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--manifest", type=Path, default=MANIFEST_PATH)
    parser.add_argument("--quarterly-output", type=Path, default=QUARTERLY_OUTPUT)
    parser.add_argument("--annual-output", type=Path, default=ANNUAL_OUTPUT)
    parser.add_argument("--national-output", type=Path, default=NATIONAL_OUTPUT)
    parser.add_argument("--schema-output", type=Path, default=SCHEMA_OUTPUT)
    parser.add_argument("--reconciliation-output", type=Path, default=RECONCILIATION_OUTPUT)
    args = parser.parse_args()

    manifest = json.loads(args.manifest.read_text(encoding="utf-8"))
    downloaded = [
        resource
        for resource in manifest.get("resources", [])
        if resource.get("status") == "downloaded"
        and resource.get("language", "en") == "en"
    ]
    quarterly_sources = [
        resource for resource in downloaded if resource["source_kind"] == "quarterly_release"
    ]
    annual_sources = [
        resource for resource in downloaded if resource["source_kind"] == "annual_release"
    ]
    if not quarterly_sources:
        raise AggregateTableError("Manifest contains no downloaded quarterly Table 14 source")
    if len(annual_sources) != 1:
        raise AggregateTableError(
            f"Expected one downloaded annual Table 14 source, found {len(annual_sources)}"
        )

    quarterly_observations = []
    annual_observations = []
    source_audit = []
    seen_source_sha = set()
    for source in quarterly_sources + annual_sources:
        source_identity = (source["source_kind"], source["sha256"])
        if source_identity in seen_source_sha:
            continue
        seen_source_sha.add(source_identity)
        observations, details = parse_source(source)
        source_audit.append(details)
        if source["source_kind"] == "quarterly_release":
            quarterly_observations.extend(observations)
        else:
            annual_observations.extend(observations)

    quarterly, revisions, revised_keys = latest_quarterly(quarterly_observations)
    annual_observations.sort(
        key=lambda value: (
            value["year"],
            value["province_territory"],
            value["program_stream"],
        )
    )
    periods = sorted({row["period"] for row in quarterly}, key=period_tuple)
    expected = period_range(FIRST_EXPECTED_PERIOD, periods[-1])
    missing_periods = [period for period in expected if period not in periods]
    annual_checks = annual_reconciliation(quarterly, annual_observations)
    national = national_rows(quarterly, revised_keys, annual_checks)

    quarterly_fields = [
        "period",
        "year",
        "quarter",
        "province_territory",
        "province_territory_raw",
        "program_stream",
        "program_stream_raw",
        "positive_lmias",
        "source_dataset_id",
        "source_resource_id",
        "source_language",
        "source_release_period",
        "source_filename",
        "source_sha256",
        "source_sheet",
        "source_row",
    ]
    annual_fields = [
        "year",
        "province_territory",
        "province_territory_raw",
        "program_stream",
        "program_stream_raw",
        "positive_lmias",
        "source_dataset_id",
        "source_resource_id",
        "source_language",
        "source_filename",
        "source_sha256",
        "source_sheet",
        "source_row",
    ]
    national_fields = [
        "period",
        "year",
        "quarter",
        "positive_lmias",
        "contributing_rows",
        "source_release_periods",
        "quality_flags",
    ]
    atomic_write(args.quarterly_output.resolve(), csv_bytes(quarterly, quarterly_fields))
    atomic_write(args.annual_output.resolve(), csv_bytes(annual_observations, annual_fields))
    atomic_write(args.national_output.resolve(), csv_bytes(national, national_fields))
    write_json(args.schema_output.resolve(), schema_document())

    reconciliation = {
        "schema_version": "0.2.0",
        "measure": "positive_lmias",
        "expected_quarterly_coverage_start": FIRST_EXPECTED_PERIOD,
        "observed_quarterly_coverage": {
            "start": periods[0],
            "end": periods[-1],
            "period_count": len(periods),
            "periods": periods,
            "missing_periods": missing_periods,
            "status": "complete" if not missing_periods else "partial",
        },
        "source_audit": source_audit,
        "quarterly_observation_count": len(quarterly),
        "annual_observation_count": len(annual_observations),
        "national_period_count": len(national),
        "overlap_revision_count": len(revisions),
        "overlap_revisions": revisions,
        "annual_reconciliation": {
            "comparison_count": len(annual_checks),
            "match_count": sum(value["status"] == "match" for value in annual_checks),
            "mismatch_count": sum(value["status"] == "mismatch" for value in annual_checks),
            "missing_annual_count": sum(
                value["status"] == "missing_annual_observation" for value in annual_checks
            ),
            "missing_quarterly_count": sum(
                value["status"] == "missing_quarterly_observation" for value in annual_checks
            ),
            "national_year_summary": annual_national_summary(annual_checks),
            "comparisons": annual_checks,
        },
        "limitations": [
            "The official quarterly catalogue is a rolling window; missing historical "
            "releases remain missing.",
            "The newest overlapping official release is selected after revisions are recorded.",
            "Quarterly and annual Table 14 do not always publish the same program-stream "
            "categories.",
            "Annual reconciliation is a check, not a substitute for missing quarterly "
            "observations.",
            "Complete all-stream negative LMIA decisions are not published in this data family.",
        ],
    }
    write_json(args.reconciliation_output.resolve(), reconciliation)
    print(
        f"Built {len(quarterly):,} quarterly observations for {len(periods)} periods "
        f"({periods[0]} to {periods[-1]}); {len(missing_periods)} expected periods missing."
    )


if __name__ == "__main__":
    main()
