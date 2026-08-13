"""Build one canonical CSV from the official LMIA employer-list resources.

The builder fails closed on an unknown schema or an unexpected row. Known source
defects are handled explicitly: the mislinked 2021 Q2 negative CSV is replaced by
the checksum-pinned official ESDC workbook, while the incorrect English measure
headers in 2022 Q2 positive are interpreted using an exact matching official French
edition. Every accepted row retains source-file, sheet, row, resource, and checksum
provenance.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import os
import re
import tempfile
import unicodedata
from collections import Counter
from contextlib import contextmanager
from dataclasses import dataclass, field, replace
from pathlib import Path
from typing import Any, Iterable, Iterator, TextIO

from lmiacsv.census_employer_resources import (
    clean_text,
    decode_csv,
    normalized_header,
    schema_signature,
)
from lmiacsv.scrape_catalogue import write_json


WORKSPACE_ROOT = Path(__file__).resolve().parents[1]
CATALOGUE_PATH = WORKSPACE_ROOT / "data" / "meta" / "employers.catalogue.json"
CENSUS_PATH = WORKSPACE_ROOT / "data" / "meta" / "employers.census.json"
RAW_MANIFEST_PATH = WORKSPACE_ROOT / "data" / "raw" / "employer-lists" / "manifest.json"
OUTPUT_PATH = WORKSPACE_ROOT / "data" / "full" / "employers.csv"
RECONCILIATION_PATH = WORKSPACE_ROOT / "data" / "meta" / "employers.reconciliation.json"
SCHEMA_PATH = WORKSPACE_ROOT / "data" / "meta" / "employers.schema.json"

NORMALIZER_VERSION = "0.3.0"
MISLINKED_2021_Q2_NEGATIVE = "4f9a23d3-c6a6-4472-94ad-ec9646351539"
RESOLVED_2022_Q2_POSITIVE_EN = "dd627095-614a-45df-a7ef-df0a4a46a812"
RESOLVED_2022_Q2_POSITIVE_FR = "d93b17b3-4c84-45d0-8d4a-b44a54f733e6"
RESOLVED_2022_Q2_QUALITY_FLAG = (
    "english_requested_headers_interpreted_as_approved_from_matching_french_release"
)
RECOVERED_2021_Q2_PATH = (
    WORKSPACE_ROOT
    / "evidence"
    / "open-government"
    / "2021-q2-negative-resource"
    / "2026-08-11"
    / "sources"
    / "TFWP_2021Q2_Negative_EN.xlsx"
)
RECOVERED_2021_Q2_SHA256 = "ee14d30f2b93b5e607974f8301ad7c09126ed73bd6a678b3edb9fcc5bc3c9f67"
RECOVERED_2021_Q2_URL = (
    "https://www.esdc.gc.ca/ouvert-open/bca-seb/imt-lmi/TFWP_2021Q2_Negative_EN.xlsx"
)

NOTE_RE = re.compile(r"^(notes?|remarques?)\s*:?$", re.I)
POSTAL_RE = re.compile(
    r"\b([ABCEGHJ-NPRSTVXY]\d[ABCEGHJ-NPRSTVWXYZ])\s?"
    r"(\d[ABCEGHJ-NPRSTVWXYZ]\d)\b",
    re.I,
)
OCCUPATION_RE = re.compile(r"^\s*(\d{4,5})\s*(?:[-–—:]\s*)?(.*)$")
NOC_VERSION_RE = re.compile(r"\b(?:NOC|CNP)\)?\s*(2011|2016|2021)\b", re.I)

CANONICAL_FIELDS = [
    "record_id",
    "period_key",
    "period_year",
    "period_quarter",
    "period_precision",
    "decision_status",
    "language",
    "comparability_regime",
    "province_raw",
    "province",
    "province_code",
    "program_stream_raw",
    "program_stream",
    "program_stream_key",
    "employer_raw",
    "employer",
    "employer_key",
    "address_raw",
    "address",
    "postal_code",
    "occupation_raw",
    "noc_version",
    "noc_code",
    "occupation_title",
    "occupation_label",
    "incorporate_status_raw",
    "incorporate_status",
    "incorporate_status_key",
    "lmia_measure",
    "lmia_count",
    "position_measure",
    "position_count",
    "source_dataset_id",
    "source_resource_id",
    "source_resource_url",
    "source_download_url",
    "source_final_url",
    "source_file",
    "source_format",
    "source_catalogue_format",
    "source_sheet",
    "source_row",
    "source_sha256",
    "source_schema_signature",
    "source_disposition",
    "quality_flags",
    "normalizer_version",
]

PROVINCES = {
    "alberta": ("Alberta", "AB"),
    "british columbia": ("British Columbia", "BC"),
    "manitoba": ("Manitoba", "MB"),
    "new brunswick": ("New Brunswick", "NB"),
    "newfoundland and labrador": ("Newfoundland and Labrador", "NL"),
    "terre neuve et labrador": ("Newfoundland and Labrador", "NL"),
    "nova scotia": ("Nova Scotia", "NS"),
    "nouvelle ecosse": ("Nova Scotia", "NS"),
    "nouvelle cosse": ("Nova Scotia", "NS"),
    "northwest territories": ("Northwest Territories", "NT"),
    "territoires du nord ouest": ("Northwest Territories", "NT"),
    "nunavut": ("Nunavut", "NU"),
    "ontario": ("Ontario", "ON"),
    "prince edward island": ("Prince Edward Island", "PE"),
    "ile du prince edouard": ("Prince Edward Island", "PE"),
    "le du prince douard": ("Prince Edward Island", "PE"),
    "quebec": ("Quebec", "QC"),
    "qu bec": ("Quebec", "QC"),
    "saskatchewan": ("Saskatchewan", "SK"),
    "yukon": ("Yukon", "YT"),
    "colombie britannique": ("British Columbia", "BC"),
    "nouveau brunswick": ("New Brunswick", "NB"),
    "outside canada": ("Outside Canada", "INT"),
    "a l exterieur du canada": ("Outside Canada", "INT"),
    "les employeurs exercant leurs activites au canada avec un siege social a l exterieur du canada": (
        "Employers carrying on business in Canada with head office outside of Canada",
        "INT",
    ),
    "les employeurs exercant leurs activites au canada avec un siege social a l exterieur du canad": (
        "Employers carrying on business in Canada with head office outside of Canada",
        "INT",
    ),
    "les employeurs exer ant leurs activit s au canada avec un si ge social l ext rieur du canada": (
        "Employers carrying on business in Canada with head office outside of Canada",
        "INT",
    ),
    "employers carrying on business in canada with head office outside of canada": (
        "Employers carrying on business in Canada with head office outside of Canada",
        "INT",
    ),
}

STREAMS = {
    "high wage": "High-wage",
    "high wage stream": "High-wage",
    "low wage": "Low-wage",
    "low wage stream": "Low-wage",
    "primary agriculture": "Primary agriculture",
    "primary agriculture stream": "Primary agriculture",
    "agricultural stream": "Agricultural stream",
    "global talent stream": "Global Talent Stream",
    "global talent": "Global Talent Stream",
    "in home caregiver": "In-home caregiver",
    "in home caregivers": "In-home caregiver",
    "permanent residence": "Permanent residence",
    "haut salaire": "High-wage",
    "bas salaire": "Low-wage",
    "agriculture primaire": "Primary agriculture",
    "volet des talents mondiaux": "Global Talent Stream",
    "talent mondial": "Global Talent Stream",
    "residence permanente uniquement": "Permanent Resident Only",
    "resident permanent seulement": "Permanent Resident Only",
    "r sident permanent seulement": "Permanent Resident Only",
    "permanent resident only": "Permanent Resident Only",
    "non specifie": "Unspecified",
    "non sp cifi": "Unspecified",
}

PROGRAM_STREAM_KEYS = {
    "high wage": "high_wage",
    "low wage": "low_wage",
    "primary agriculture": "primary_agriculture",
    "agricultural stream": "agricultural_stream",
    "global talent stream": "global_talent_stream",
    "in home caregiver": "in_home_caregiver",
    "permanent residence": "permanent_residence",
    "permanent resident only": "permanent_resident_only",
    "unspecified": "unspecified",
    "99": "source_code_99",
}

HEADER_ALIASES = {
    "province territory": "province",
    "province territoire": "province",
    "stream": "program_stream",
    "program stream": "program_stream",
    "volet": "program_stream",
    "volet du programme": "program_stream",
    "employer": "employer",
    "employer name": "employer",
    "employeur": "employer",
    "address": "address",
    "adresse": "address",
    "occupation": "occupation",
    "profession": "occupation",
    "profession sous cnp 2011": "occupation",
    "occupations under noc 2011": "occupation",
    "occupaitons under noc 2011": "occupation",
    "incorporate status": "incorporate_status",
    "incorporer le statut": "incorporate_status",
    "approved lmias": "lmia_count",
    "requested lmias": "lmia_count",
    "eimt approuves": "lmia_count",
    "eimt demandes": "lmia_count",
    "positions": "position_count",
    "positions approved": "position_count",
    "approved positions": "position_count",
    "positions requested": "position_count",
    "requested positions": "position_count",
    "postes approuves": "position_count",
    "postes demandes": "position_count",
}


class CanonicalBuildError(RuntimeError):
    """Raised when an unmodelled source condition would make output dishonest."""


@dataclass(frozen=True)
class SourcePlan:
    catalogue: dict[str, Any]
    census: dict[str, Any]
    raw: dict[str, Any]
    path: Path
    format: str
    catalogue_format: str
    file_name: str
    sheet: str
    header_row: int
    header_values: list[str]
    first_note_row: int | None
    sha256: str
    final_url: str
    download_url: str
    schema_signature: str
    disposition: str
    anomaly_flags: tuple[str, ...] = ()
    interpretation_evidence: dict[str, Any] = field(default_factory=dict)

    @property
    def resource_id(self) -> str:
        return self.catalogue["resource_id"]

    @property
    def decision(self) -> str:
        return self.catalogue["decision"]


@dataclass
class ResourceStats:
    records: int = 0
    position_count: int = 0
    lmia_count: int = 0
    records_with_lmia_count: int = 0
    zero_position_records: int = 0
    blank_province_records: int = 0
    blank_program_stream_records: int = 0
    blank_occupation_records: int = 0
    quality_flags: Counter[str] = field(default_factory=Counter)

    def observe(self, record: dict[str, Any]) -> None:
        self.records += 1
        self.position_count += int(record["position_count"])
        if record["lmia_count"] != "":
            self.lmia_count += int(record["lmia_count"])
            self.records_with_lmia_count += 1
        self.zero_position_records += record["position_count"] == 0
        self.blank_province_records += not bool(record["province"])
        self.blank_program_stream_records += not bool(record["program_stream"])
        self.blank_occupation_records += not bool(record["occupation_raw"])
        for flag in str(record["quality_flags"]).split("|"):
            if flag:
                self.quality_flags[flag] += 1


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def raw_text(value: Any) -> str:
    if value is None:
        return ""
    return unicodedata.normalize("NFC", str(value))


def search_key(value: str) -> str:
    folded = unicodedata.normalize("NFKD", value.casefold())
    characters = (
        character if character.isalnum() else " "
        for character in folded
        if not unicodedata.combining(character)
    )
    return " ".join("".join(characters).split())


def integer(value: Any, label: str, resource_id: str, row_number: int) -> int:
    if isinstance(value, bool):
        raise CanonicalBuildError(f"{resource_id} row {row_number}: {label} cannot be boolean")
    if isinstance(value, int):
        parsed = value
    elif isinstance(value, float) and value.is_integer():
        parsed = int(value)
    else:
        text = clean_text(value).replace(",", "")
        if not text.isdigit():
            raise CanonicalBuildError(
                f"{resource_id} row {row_number}: {label} is not an integer: {value!r}"
            )
        parsed = int(text)
    if parsed < 0:
        raise CanonicalBuildError(
            f"{resource_id} row {row_number}: {label} cannot be negative: {parsed}"
        )
    return parsed


def stable_record_id(plan: SourcePlan, row_number: int) -> str:
    identity = f"{plan.resource_id}:{plan.sha256}:{plan.sheet}:{row_number}".encode()
    return f"lmia-{hashlib.sha256(identity).hexdigest()[:24]}"


def normalize_province(value: str) -> tuple[str, str]:
    if not value:
        return "", ""
    return PROVINCES.get(search_key(value), (clean_text(value), ""))


def normalize_stream(value: str) -> str:
    if not value:
        return ""
    cleaned = clean_text(value)
    return STREAMS.get(search_key(cleaned), cleaned)


def program_stream_key(value: str) -> str:
    if not value:
        return ""
    key = search_key(value)
    try:
        return PROGRAM_STREAM_KEYS[key]
    except KeyError as error:
        raise CanonicalBuildError(
            f"Program stream has no language-neutral key: {value!r}"
        ) from error


def normalize_incorporate_status(value: str) -> str:
    cleaned = clean_text(value)
    key = search_key(cleaned)
    if key in {"incorporated", "incorpore", "incorpor", "yes"}:
        return "Incorporated"
    if key in {"unincorporated", "not incorporated", "no"}:
        return "Unincorporated"
    if key in {"unknown", "inconnu", "inconnue"}:
        return "Unknown"
    if key in {"partnership", "partenariat"}:
        return "Partnership"
    if key in {
        "sole proprietorship",
        "entreprise a proprietaire unique",
        "entreprise proprietaire unique",
        "entreprise propri taire unique",
    }:
        return "Sole Proprietorship"
    if key in {"co operative", "cooperative", "coop rative"}:
        return "Co-operative"
    if key in {"non profit organization", "organisme sans but lucratif"}:
        return "Non-Profit Organization"
    if key in {
        "registered charity",
        "organisme de bienfaisance enregistre",
        "organisme de bienfaisance enregist",
    }:
        return "Registered Charity"
    if key in {"no response provided", "aucune reponse fournie"}:
        return "No response provided"
    return cleaned


def incorporate_status_key(value: str) -> str:
    return search_key(value).replace(" ", "_") if value else ""


def postal_code(address: str) -> str:
    match = POSTAL_RE.search(address)
    if not match:
        return ""
    return f"{match.group(1).upper()} {match.group(2).upper()}"


def source_noc_version(plan: SourcePlan, code: str) -> str:
    title = plan.catalogue.get("title", {})
    title_text = " ".join(str(value) for value in title.values()) if isinstance(title, dict) else str(title)
    evidence = " ".join([title_text, *plan.header_values])
    versions = set(NOC_VERSION_RE.findall(evidence))
    if len(code) == 5:
        if versions and versions != {"2021"}:
            raise CanonicalBuildError(
                f"{plan.resource_id}: five-digit NOC code conflicts with source metadata "
                f"{sorted(versions)!r}"
            )
        return "2021"
    if len(versions) != 1:
        raise CanonicalBuildError(
            f"{plan.resource_id}: four-digit NOC code has ambiguous source version "
            f"{sorted(versions)!r}"
        )
    version = versions.pop()
    if version == "2021":
        raise CanonicalBuildError(
            f"{plan.resource_id}: four-digit NOC code conflicts with NOC 2021 metadata"
        )
    return version


def parse_occupation(value: str, plan: SourcePlan) -> tuple[str, str, str, str]:
    cleaned = clean_text(value)
    if not cleaned:
        return "", "", "", ""
    match = OCCUPATION_RE.match(cleaned)
    if not match:
        return "", "", cleaned, cleaned
    code = match.group(1)
    title = match.group(2).strip()
    version = source_noc_version(plan, code)
    label = f"{code} — {title}" if title else code
    return version, code, title, label


def comparability_regime(period: dict[str, Any]) -> str:
    year = int(period["year"])
    quarter = period.get("quarter")
    if year < 2018:
        return "legacy_pre_2018"
    if year < 2023 or (year == 2023 and quarter in {"Q1", "Q2", "Q3"}):
        return "pr_only_lmias_excluded"
    return "pr_only_lmias_included"


def header_roles(header_values: Iterable[Any]) -> list[str]:
    roles = []
    for value in header_values:
        normalized = normalized_header(value)
        role = HEADER_ALIASES.get(normalized)
        if role is None:
            raise CanonicalBuildError(f"Unrecognized source header: {clean_text(value)!r}")
        roles.append(role)
    if len(roles) != len(set(roles)):
        raise CanonicalBuildError(f"Source header maps two columns to one role: {roles!r}")
    required = {"employer", "address", "position_count"}
    missing = sorted(required - set(roles))
    if missing:
        raise CanonicalBuildError(f"Source schema is missing required fields: {missing}")
    return roles


def header_decision_signal(header_values: Iterable[Any]) -> str | None:
    values = [normalized_header(value) for value in header_values]
    if any("approved" in value or "approuve" in value for value in values):
        return "positive"
    if any("requested" in value or "demande" in value for value in values):
        return "negative"
    return None


def is_header_row(values: Iterable[Any]) -> bool:
    normalized = [normalized_header(value) for value in values]
    return any(value in normalized for value in ("employer", "employer name", "employeur"))


def source_path(path_value: str) -> Path:
    path = Path(path_value)
    return path if path.is_absolute() else WORKSPACE_ROOT / path


def index_resources(document: dict[str, Any]) -> dict[str, dict[str, Any]]:
    return {
        resource["resource_id"]: resource
        for resource in document.get("resources", [])
        if resource.get("resource_id") and resource.get("language") == "en"
    }


def resource_by_id(document: dict[str, Any], resource_id: str) -> dict[str, Any]:
    matches = [
        resource
        for resource in document.get("resources", [])
        if resource.get("resource_id") == resource_id
    ]
    if len(matches) != 1:
        raise CanonicalBuildError(
            f"Expected exactly one resource {resource_id}, found {len(matches)}"
        )
    return matches[0]


def validate_input_sets(
    catalogue: dict[str, dict[str, Any]],
    census: dict[str, dict[str, Any]],
    raw: dict[str, dict[str, Any]],
) -> None:
    sets = {"catalogue": set(catalogue), "census": set(census), "raw manifest": set(raw)}
    if len({frozenset(value) for value in sets.values()}) == 1:
        return
    universe = set().union(*sets.values())
    details = []
    for name, values in sets.items():
        missing = sorted(universe - values)
        if missing:
            details.append(f"{name} missing {missing}")
    raise CanonicalBuildError("Input resource sets disagree: " + "; ".join(details))


def selected_sheet(census: dict[str, Any]) -> dict[str, Any]:
    name = census.get("selected_sheet")
    matches = [sheet for sheet in census.get("sheets", []) if sheet.get("name") == name]
    if len(matches) != 1 or not matches[0].get("header"):
        raise CanonicalBuildError(
            f"{census['resource_id']}: census does not identify exactly one header-bearing sheet"
        )
    return matches[0]


def catalogued_source_plan(
    catalogue: dict[str, Any],
    census: dict[str, Any],
    raw: dict[str, Any],
    *,
    allow_decision_mismatch: bool = False,
) -> SourcePlan:
    resource_id = catalogue["resource_id"]
    if raw.get("status") != "downloaded":
        raise CanonicalBuildError(f"{resource_id}: raw resource was not downloaded")
    if raw.get("sha256") != census.get("sha256"):
        raise CanonicalBuildError(f"{resource_id}: census and raw checksums disagree")
    path = source_path(raw["archive_path"])
    if not path.is_file():
        raise CanonicalBuildError(f"{resource_id}: archived source is missing: {path}")
    if file_sha256(path) != raw["sha256"]:
        raise CanonicalBuildError(f"{resource_id}: archived source checksum mismatch")

    sheet = selected_sheet(census)
    headers = [clean_text(value) for value in sheet["header"]["values"]]
    roles = header_roles(headers)
    signal = header_decision_signal(headers)
    if signal is not None and signal != catalogue["decision"] and not allow_decision_mismatch:
        raise CanonicalBuildError(
            f"{resource_id}: {catalogue['decision']} catalogue resource has {signal} headers"
        )
    if len(roles) not in {3, 5, 6, 8}:
        raise CanonicalBuildError(f"{resource_id}: unsupported {len(roles)}-column schema")

    anomalies = []
    if not raw.get("format_matches_catalogue", True):
        anomalies.append(
            "catalogue_format_mismatch_"
            f"{raw.get('catalogue_format', catalogue['format']).lower()}_to_"
            f"{raw['format'].lower()}"
        )
    if signal is None:
        anomalies.append("measure_semantics_inferred_from_decision_title")
    elif signal != catalogue["decision"]:
        anomalies.append("decision_header_mismatch_preserved_for_bilingual_audit")
    return SourcePlan(
        catalogue=catalogue,
        census=census,
        raw=raw,
        path=path,
        format=raw["format"],
        catalogue_format=raw.get("catalogue_format", catalogue["format"]),
        file_name=raw["source_filename"],
        sheet=sheet["name"],
        header_row=sheet["header"]["row"],
        header_values=headers,
        first_note_row=sheet.get("first_note_row"),
        sha256=raw["sha256"],
        final_url=raw.get("final_url", catalogue["download_url"]),
        download_url=catalogue["download_url"],
        schema_signature=census["selected_schema_signature"],
        disposition="normalized",
        anomaly_flags=tuple(anomalies),
    )


def recovered_2021_q2_plan(
    catalogue: dict[str, Any], census: dict[str, Any], raw: dict[str, Any]
) -> SourcePlan:
    path = RECOVERED_2021_Q2_PATH
    if not path.is_file():
        raise CanonicalBuildError(
            "The recovered official 2021 Q2 negative workbook is missing. "
            f"Run lmiacsv receipts or restore {path}."
        )
    digest = file_sha256(path)
    if digest != RECOVERED_2021_Q2_SHA256:
        raise CanonicalBuildError(
            f"Recovered 2021 Q2 checksum mismatch: expected "
            f"{RECOVERED_2021_Q2_SHA256}, got {digest}"
        )
    headers = [
        "Province/Territory",
        "Program Stream",
        "Employer",
        "Address",
        "Occupation",
        "Requested Positions",
    ]
    header = {"normalized_values": [normalized_header(value) for value in headers]}
    return SourcePlan(
        catalogue=catalogue,
        census=census,
        raw=raw,
        path=path,
        format="XLSX",
        catalogue_format=raw.get("catalogue_format", catalogue["format"]),
        file_name=path.name,
        sheet="Negative_EN",
        header_row=2,
        header_values=headers,
        first_note_row=None,
        sha256=digest,
        final_url=RECOVERED_2021_Q2_URL,
        download_url=RECOVERED_2021_Q2_URL,
        schema_signature=schema_signature(header),
        disposition="recovered_official_alternate",
        anomaly_flags=(
            "catalogue_resource_mislinked_to_2020_q2",
            "recovered_from_official_esdc_xlsx",
        ),
    )


def normalized_record_signature(
    plan: SourcePlan,
) -> tuple[Counter[tuple[Any, ...]], dict[str, int]]:
    signature: Counter[tuple[Any, ...]] = Counter()
    totals = {"records": 0, "lmias": 0, "positions": 0}
    with source_rows(plan) as rows:
        for record in canonical_records(plan, rows):
            identity = (
                record["province_code"],
                record["program_stream_key"],
                record["postal_code"],
                record["noc_version"],
                record["noc_code"],
                record["lmia_count"],
                record["position_count"],
                record["incorporate_status_key"],
                record["employer_key"],
            )
            signature[identity] += 1
            totals["records"] += 1
            totals["lmias"] += int(record["lmia_count"] or 0)
            totals["positions"] += int(record["position_count"])
    return signature, totals


def resolved_2022_q2_positive_plan(
    catalogue_document: dict[str, Any],
    census_document: dict[str, Any],
    raw_document: dict[str, Any],
) -> SourcePlan:
    en_catalogue = resource_by_id(catalogue_document, RESOLVED_2022_Q2_POSITIVE_EN)
    fr_catalogue = resource_by_id(catalogue_document, RESOLVED_2022_Q2_POSITIVE_FR)
    en_census = resource_by_id(census_document, RESOLVED_2022_Q2_POSITIVE_EN)
    fr_census = resource_by_id(census_document, RESOLVED_2022_Q2_POSITIVE_FR)
    en_raw = resource_by_id(raw_document, RESOLVED_2022_Q2_POSITIVE_EN)
    fr_raw = resource_by_id(raw_document, RESOLVED_2022_Q2_POSITIVE_FR)

    expected_identity = ("2022-q2", "positive")
    for language, resource in (("en", en_catalogue), ("fr", fr_catalogue)):
        identity = (resource["period"]["key"], resource["decision"])
        if resource.get("language") != language or identity != expected_identity:
            raise CanonicalBuildError(
                f"{resource['resource_id']}: unexpected bilingual override identity "
                f"{resource.get('language')!r}, {identity!r}"
            )

    en_plan = catalogued_source_plan(
        en_catalogue,
        en_census,
        en_raw,
        allow_decision_mismatch=True,
    )
    fr_plan = catalogued_source_plan(fr_catalogue, fr_census, fr_raw)
    en_signal = header_decision_signal(en_plan.header_values)
    fr_signal = header_decision_signal(fr_plan.header_values)
    if en_signal != "negative" or fr_signal != "positive":
        raise CanonicalBuildError(
            "2022 Q2 positive bilingual override no longer has requested English "
            f"headers and approved French headers: en={en_signal!r}, fr={fr_signal!r}"
        )

    en_signature, en_totals = normalized_record_signature(en_plan)
    fr_signature, fr_totals = normalized_record_signature(fr_plan)
    if en_signature != fr_signature or en_totals != fr_totals:
        raise CanonicalBuildError(
            "2022 Q2 positive English/French records no longer match exactly; "
            "refusing the measure-header override"
        )

    anomaly_flags = tuple(
        flag
        for flag in en_plan.anomaly_flags
        if flag != "decision_header_mismatch_preserved_for_bilingual_audit"
    ) + (RESOLVED_2022_Q2_QUALITY_FLAG,)
    return replace(
        en_plan,
        disposition="normalized_with_bilingual_measure_override",
        anomaly_flags=anomaly_flags,
        interpretation_evidence={
            "reason": (
                "The English positive workbook labels its measures Requested LMIAs and "
                "Requested Positions. Its paired official French positive workbook labels "
                "the same measures approved and matches every language-neutral normalized "
                "record and measure."
            ),
            "english_header_signal": en_signal,
            "french_header_signal": fr_signal,
            "paired_french_resource_id": fr_plan.resource_id,
            "paired_french_source_file": fr_plan.file_name,
            "paired_french_sha256": fr_plan.sha256,
            "normalized_record_multiset_match": True,
            "totals_match": True,
            "verified_totals": en_totals,
        },
    )


def build_source_plans(
    catalogue_document: dict[str, Any],
    census_document: dict[str, Any],
    raw_document: dict[str, Any],
) -> tuple[list[SourcePlan], list[dict[str, Any]]]:
    catalogue = index_resources(catalogue_document)
    census = index_resources(census_document)
    raw = index_resources(raw_document)
    validate_input_sets(catalogue, census, raw)

    plans = []
    quarantines = []
    ordered = sorted(
        catalogue.values(),
        key=lambda resource: (
            resource["period"]["year"],
            resource["period"].get("quarter") or "",
            resource["decision"],
        ),
    )
    for resource in ordered:
        resource_id = resource["resource_id"]
        if resource_id == RESOLVED_2022_Q2_POSITIVE_EN:
            plan = resolved_2022_q2_positive_plan(
                catalogue_document,
                census_document,
                raw_document,
            )
        elif resource_id == MISLINKED_2021_Q2_NEGATIVE:
            plan = recovered_2021_q2_plan(resource, census[resource_id], raw[resource_id])
        else:
            plan = catalogued_source_plan(resource, census[resource_id], raw[resource_id])
        plans.append(plan)
    return plans, quarantines


@contextmanager
def source_rows(plan: SourcePlan) -> Iterator[Iterator[tuple[int, list[Any]]]]:
    if plan.format == "CSV":
        text, _encoding = decode_csv(plan.path)
        sample = text[:8192]
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
        reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
        yield ((row_number, list(values)) for row_number, values in enumerate(reader, 1))
        return
    if plan.format != "XLSX":
        raise CanonicalBuildError(
            f"{plan.resource_id}: unsupported detected format {plan.format!r}"
        )
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError as error:
        raise RuntimeError("openpyxl 3.1.5 is required to build canonical records") from error
    workbook = load_workbook(plan.path, read_only=True, data_only=True)
    try:
        if plan.sheet not in workbook.sheetnames:
            raise CanonicalBuildError(
                f"{plan.resource_id}: expected sheet {plan.sheet!r}, found {workbook.sheetnames!r}"
            )
        sheet = workbook[plan.sheet]
        rows = sheet.iter_rows(values_only=True)
        yield ((row_number, list(values)) for row_number, values in enumerate(rows, 1))
    finally:
        workbook.close()


def row_mapping(roles: list[str], values: list[Any], resource_id: str, row: int) -> dict[str, Any]:
    width = len(roles)
    extras = values[width:]
    if any(clean_text(value) for value in extras):
        raise CanonicalBuildError(
            f"{resource_id} row {row}: nonempty cells beyond the {width}-column schema"
        )
    padded = values[:width] + [None] * max(0, width - len(values))
    return dict(zip(roles, padded, strict=True))


def canonical_records(
    plan: SourcePlan, rows: Iterable[tuple[int, list[Any]]]
) -> Iterator[dict[str, Any]]:
    roles = header_roles(plan.header_values)
    period = plan.catalogue["period"]
    current_province = ""
    current_stream = ""
    current_employer = ""
    current_address = ""
    for row_number, values in rows:
        if row_number < plan.header_row and "province" not in roles:
            leading = [clean_text(value) for value in values]
            while leading and not leading[-1]:
                leading.pop()
            if len(leading) == 1:
                province, code = normalize_province(leading[0])
                if code:
                    current_province = province
            continue
        if row_number <= plan.header_row:
            continue
        if plan.first_note_row is not None and row_number >= plan.first_note_row:
            break
        trimmed = list(values)
        while trimmed and not clean_text(trimmed[-1]):
            trimmed.pop()
        if not trimmed:
            continue
        first = clean_text(trimmed[0])
        if NOTE_RE.match(first):
            break
        if is_header_row(trimmed):
            continue

        # The 2014 files use single-cell province headings and repeat the header
        # beneath each heading instead of supplying a province column per record.
        if "province" not in roles:
            province, code = normalize_province(first)
            repeated_heading = len(trimmed) == 2 and search_key(
                clean_text(trimmed[1])
            ) == search_key(first)
            if code and (len(trimmed) == 1 or repeated_heading):
                current_province = province
                current_stream = ""
                current_employer = ""
                current_address = ""
                continue

        mapped = row_mapping(roles, trimmed, plan.resource_id, row_number)
        province_raw = raw_text(mapped.get("province"))
        province_cell = clean_text(mapped.get("province"))
        if province_cell:
            province, _code = normalize_province(province_cell)
            if province != current_province:
                current_stream = ""
                current_employer = ""
                current_address = ""
            current_province = province
        stream_raw = raw_text(mapped.get("program_stream"))
        stream_cell = clean_text(mapped.get("program_stream"))
        if stream_cell:
            normalized_stream = normalize_stream(stream_cell)
            if normalized_stream != current_stream:
                current_employer = ""
                current_address = ""
            current_stream = normalized_stream

        employer_raw = raw_text(mapped.get("employer"))
        employer_cell = clean_text(mapped.get("employer"))
        if employer_cell:
            current_employer = employer_cell
        employer = current_employer
        if not employer:
            raise CanonicalBuildError(
                f"{plan.resource_id} row {row_number}: record has no employer"
            )
        address_raw = raw_text(mapped.get("address"))
        address_cell = clean_text(mapped.get("address"))
        if address_cell:
            current_address = address_cell
        address = current_address
        occupation_raw = raw_text(mapped.get("occupation"))
        occupation = clean_text(mapped.get("occupation"))
        incorporate_raw = raw_text(mapped.get("incorporate_status"))
        incorporate = normalize_incorporate_status(incorporate_raw)
        province, province_code = normalize_province(current_province)
        noc_version, noc_code, occupation_title, occupation_label = parse_occupation(
            occupation, plan
        )
        lmia_value = mapped.get("lmia_count")
        lmia_count: int | str = ""
        if "lmia_count" in mapped:
            lmia_count = integer(lmia_value, "LMIA count", plan.resource_id, row_number)
        position_count = integer(
            mapped.get("position_count"), "position count", plan.resource_id, row_number
        )
        quality_flags = list(plan.anomaly_flags)
        if not province:
            quality_flags.append("missing_province")
        elif not province_code:
            quality_flags.append("unresolved_province")
        if occupation and not noc_code:
            quality_flags.append("unparsed_occupation_code")
        if position_count == 0:
            quality_flags.append("zero_positions_preserved")
        if not employer_cell:
            quality_flags.append("employer_forward_filled_from_grouped_source")
        if not address_cell:
            quality_flags.append("address_forward_filled_from_grouped_source")

        decision = plan.decision
        yield {
            "record_id": stable_record_id(plan, row_number),
            "period_key": period["key"],
            "period_year": period["year"],
            "period_quarter": period.get("quarter") or "",
            "period_precision": period["precision"],
            "decision_status": decision,
            "language": plan.catalogue["language"],
            "comparability_regime": comparability_regime(period),
            "province_raw": province_raw,
            "province": province,
            "province_code": province_code,
            "program_stream_raw": stream_raw,
            "program_stream": current_stream,
            "program_stream_key": program_stream_key(current_stream),
            "employer_raw": employer_raw,
            "employer": employer,
            "employer_key": search_key(employer),
            "address_raw": address_raw,
            "address": address,
            "postal_code": postal_code(address),
            "occupation_raw": occupation_raw,
            "noc_version": noc_version,
            "noc_code": noc_code,
            "occupation_title": occupation_title,
            "occupation_label": occupation_label,
            "incorporate_status_raw": incorporate_raw,
            "incorporate_status": incorporate,
            "incorporate_status_key": incorporate_status_key(incorporate),
            "lmia_measure": (
                (
                    "approved_lmias"
                    if decision == "positive"
                    else "requested_lmias_on_negative_decisions"
                )
                if "lmia_count" in mapped
                else ""
            ),
            "lmia_count": lmia_count,
            "position_measure": (
                "approved_positions"
                if decision == "positive"
                else "requested_positions_on_negative_decisions"
            ),
            "position_count": position_count,
            "source_dataset_id": plan.catalogue["dataset_id"],
            "source_resource_id": plan.resource_id,
            "source_resource_url": plan.catalogue["resource_page_url"],
            "source_download_url": plan.download_url,
            "source_final_url": plan.final_url,
            "source_file": plan.file_name,
            "source_format": plan.format,
            "source_catalogue_format": plan.catalogue_format,
            "source_sheet": plan.sheet,
            "source_row": row_number,
            "source_sha256": plan.sha256,
            "source_schema_signature": plan.schema_signature,
            "source_disposition": plan.disposition,
            "quality_flags": "|".join(dict.fromkeys(quality_flags)),
            "normalizer_version": NORMALIZER_VERSION,
        }


def resource_ledger_entry(plan: SourcePlan, stats: ResourceStats) -> dict[str, Any]:
    sheet = selected_sheet(plan.census)
    entry = {
        "dataset_id": plan.catalogue["dataset_id"],
        "resource_id": plan.resource_id,
        "period": plan.catalogue["period"],
        "decision": plan.decision,
        "language": plan.catalogue["language"],
        "disposition": plan.disposition,
        "resource_page_url": plan.catalogue["resource_page_url"],
        "catalogue_download_url": plan.catalogue["download_url"],
        "selected_download_url": plan.download_url,
        "catalogue_sha256": plan.raw["sha256"],
        "selected_sha256": plan.sha256,
        "catalogue_format": plan.catalogue_format,
        "detected_format": plan.format,
        "source_file": plan.file_name,
        "source_sheet": plan.sheet,
        "header_row": plan.header_row,
        "header_values": plan.header_values,
        "schema_signature": plan.schema_signature,
        "census_estimated_rows": sheet.get("estimated_data_rows"),
        "canonical_minus_census_estimate": (
            stats.records - sheet["estimated_data_rows"]
            if sheet.get("estimated_data_rows") is not None
            else None
        ),
        "canonical_records": stats.records,
        "position_count": stats.position_count,
        "records_with_lmia_count": stats.records_with_lmia_count,
        "lmia_count": stats.lmia_count if stats.records_with_lmia_count else None,
        "quality_counts": {
            "zero_position_records": stats.zero_position_records,
            "blank_province_records": stats.blank_province_records,
            "blank_program_stream_records": stats.blank_program_stream_records,
            "blank_occupation_records": stats.blank_occupation_records,
            "flags": dict(sorted(stats.quality_flags.items())),
        },
        "anomaly_flags": list(plan.anomaly_flags),
    }
    if plan.disposition == "recovered_official_alternate":
        entry["replaced_catalogue_blob"] = {
            "sha256": plan.raw["sha256"],
            "source_file": plan.raw["source_filename"],
            "reason": "byte-identical_to_catalogued_2020_q2_negative_resource",
        }
    if plan.interpretation_evidence:
        entry["interpretation_evidence"] = plan.interpretation_evidence
    return entry


def build_schema_document() -> dict[str, Any]:
    definitions = {
        "record_id": "Stable content-and-source-row identifier.",
        "period_key": "Catalogue period key; annual and multi-quarter periods stay unsplit.",
        "period_year": "Calendar year named by the source release.",
        "period_quarter": "Q1-Q4, Q1-Q2, or blank for an annual/partial-year source.",
        "period_precision": "quarter, multi_quarter, or year_or_partial_year.",
        "decision_status": "Catalogue decision list: positive or negative.",
        "language": "Language of the selected source; this build selects English.",
        "comparability_regime": "Scope regime for permanent-residence-only LMIAs.",
        "province_raw": "Literal province cell; blank where the source uses grouping.",
        "province": "Cleaned and forward-filled province or territory name.",
        "province_code": "Canadian postal abbreviation, or INT for outside Canada.",
        "program_stream_raw": "Literal stream cell; blank where the source uses grouping.",
        "program_stream": "Cleaned and forward-filled program stream when supplied.",
        "program_stream_key": "Stable language-neutral key joined to translations.csv.",
        "employer_raw": "Literal employer cell.",
        "employer": "Unicode- and whitespace-cleaned employer label.",
        "employer_key": "Case- and punctuation-folded employer search key; not entity resolution.",
        "address_raw": "Literal address cell.",
        "address": "Unicode- and whitespace-cleaned address, forward-filled when grouped.",
        "postal_code": "Canadian postal code parsed from the address when present.",
        "occupation_raw": "Literal occupation cell.",
        "noc_version": "NOC version read from source metadata and validated against code width.",
        "noc_code": "Four- or five-digit NOC code parsed from the occupation label.",
        "occupation_title": "Occupation title with the parsed NOC prefix removed.",
        "occupation_label": "Canonical NOC code and occupation title display label.",
        "incorporate_status_raw": "Literal incorporate-status cell when supplied.",
        "incorporate_status": "Cleaned incorporate status when supplied.",
        "incorporate_status_key": "Stable language-neutral key joined to translations.csv.",
        "lmia_measure": "Semantic label for lmia_count; blank when the source has no LMIA count.",
        "lmia_count": "LMIA count when the source supplies it; blank before 2021 Q4.",
        "position_measure": "Semantic label for position_count.",
        "position_count": "Published position count, including preserved zeroes.",
        "source_dataset_id": "Open Government dataset UUID.",
        "source_resource_id": "Open Government distribution/resource UUID.",
        "source_resource_url": "Human-facing Open Government resource page.",
        "source_download_url": "Selected source download URL; recovered sources use the ESDC URL.",
        "source_final_url": "Final URL after the archived download redirect, without query strings.",
        "source_file": "Government source filename retained as evidence, not identity.",
        "source_format": "Format detected from the selected file content.",
        "source_catalogue_format": "Format declared by the Open Government catalogue.",
        "source_sheet": "Selected workbook sheet, or CSV for a text source.",
        "source_row": "One-based row number in the selected source sheet or CSV.",
        "source_sha256": "SHA-256 checksum of the exact selected source file.",
        "source_schema_signature": "Stable signature of the normalized source headers.",
        "source_disposition": (
            "normalized, recovered_official_alternate, or "
            "normalized_with_bilingual_measure_override."
        ),
        "quality_flags": "Pipe-delimited row/source caveats carried into the canonical table.",
        "normalizer_version": "Version of this canonical normalization contract.",
    }
    integer_fields = {"period_year", "lmia_count", "position_count", "source_row"}
    nullable_fields = {
        "period_quarter",
        "province_raw",
        "province",
        "province_code",
        "program_stream_raw",
        "program_stream",
        "program_stream_key",
        "address_raw",
        "address",
        "postal_code",
        "occupation_raw",
        "noc_version",
        "noc_code",
        "occupation_title",
        "occupation_label",
        "incorporate_status_raw",
        "incorporate_status",
        "incorporate_status_key",
        "lmia_measure",
        "lmia_count",
        "quality_flags",
    }
    return {
        "schema_version": "0.3.0",
        "normalizer_version": NORMALIZER_VERSION,
        "format": "CSV UTF-8 with header",
        "primary_key": ["record_id"],
        "fields": [
            {
                "name": name,
                "type": "integer" if name in integer_fields else "string",
                "nullable": name in nullable_fields,
                "description": definitions[name],
            }
            for name in CANONICAL_FIELDS
        ],
        "measure_semantics": {
            "approved_lmias": "Counts from positive-decision lists where published.",
            "requested_lmias_on_negative_decisions": (
                "Requested LMIA counts attached to negative-decision list rows where published."
            ),
            "approved_positions": "Approved positions on positive-decision list rows.",
            "requested_positions_on_negative_decisions": (
                "Requested positions attached to negative-decision list rows."
            ),
        },
        "interpretation_rules": {
            "2022_q2_positive_english_measure_headers": {
                "source_resource_id": RESOLVED_2022_Q2_POSITIVE_EN,
                "paired_french_resource_id": RESOLVED_2022_Q2_POSITIVE_FR,
                "source_header_labels": ["Requested LMIAs", "Requested Positions"],
                "canonical_measure_labels": ["approved_lmias", "approved_positions"],
                "basis": (
                    "The paired official French positive workbook labels the measures "
                    "approved and must match the English workbook at language-neutral "
                    "normalized-record and measure-total levels. The build fails closed "
                    "if that evidence changes."
                ),
                "row_quality_flag": RESOLVED_2022_Q2_QUALITY_FLAG,
            }
        },
        "important_limit": (
            "Employer-list rows are not application-level observations. The table must not be used "
            "to calculate an LMIA application approval or denial rate."
        ),
    }


def open_temporary_csv(output_path: Path) -> tuple[Path, TextIO, csv.DictWriter]:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = tempfile.NamedTemporaryFile(
        mode="w",
        encoding="utf-8",
        newline="",
        prefix=".lmia-canonical-",
        suffix=".csv",
        dir=output_path.parent,
        delete=False,
    )
    writer = csv.DictWriter(temporary, fieldnames=CANONICAL_FIELDS, extrasaction="raise")
    writer.writeheader()
    return Path(temporary.name), temporary, writer


def display_path(path: Path) -> str:
    try:
        return path.resolve().relative_to(WORKSPACE_ROOT).as_posix()
    except ValueError:
        return str(path.resolve())


def build(
    catalogue_path: Path = CATALOGUE_PATH,
    census_path: Path = CENSUS_PATH,
    raw_manifest_path: Path = RAW_MANIFEST_PATH,
    output_path: Path = OUTPUT_PATH,
    reconciliation_path: Path = RECONCILIATION_PATH,
    schema_path: Path = SCHEMA_PATH,
    fail_on_quarantine: bool = False,
) -> dict[str, Any]:
    catalogue_document = json.loads(catalogue_path.read_text(encoding="utf-8"))
    census_document = json.loads(census_path.read_text(encoding="utf-8"))
    raw_document = json.loads(raw_manifest_path.read_text(encoding="utf-8"))
    plans, quarantines = build_source_plans(catalogue_document, census_document, raw_document)
    if fail_on_quarantine and quarantines:
        ids = ", ".join(resource["resource_id"] for resource in quarantines)
        raise CanonicalBuildError(f"Known quarantined resources present: {ids}")

    temporary_path, handle, writer = open_temporary_csv(output_path)
    resource_entries = []
    record_ids: set[str] = set()
    total_records = 0
    totals_by_decision: dict[str, dict[str, int]] = {
        "positive": {"records": 0, "positions": 0, "lmias": 0},
        "negative": {"records": 0, "positions": 0, "lmias": 0},
    }
    try:
        for index, plan in enumerate(plans, start=1):
            stats = ResourceStats()
            with source_rows(plan) as rows:
                for record in canonical_records(plan, rows):
                    record_id = record["record_id"]
                    if record_id in record_ids:
                        raise CanonicalBuildError(f"Duplicate canonical record ID: {record_id}")
                    record_ids.add(record_id)
                    writer.writerow(record)
                    stats.observe(record)
            if plan.resource_id == MISLINKED_2021_Q2_NEGATIVE and stats.records != 498:
                raise CanonicalBuildError(
                    f"Recovered 2021 Q2 workbook produced {stats.records} rows; expected 498"
                )
            resource_entries.append(resource_ledger_entry(plan, stats))
            total_records += stats.records
            bucket = totals_by_decision[plan.decision]
            bucket["records"] += stats.records
            bucket["positions"] += stats.position_count
            bucket["lmias"] += stats.lmia_count
            print(
                f"[{index}/{len(plans)}] {plan.catalogue['period']['key']} "
                f"{plan.decision}: {stats.records:,} rows ({plan.disposition})",
                flush=True,
            )
        handle.flush()
        os.fsync(handle.fileno())
        handle.close()
        temporary_path.chmod(0o644)
        temporary_path.replace(output_path)
    except Exception:
        handle.close()
        temporary_path.unlink(missing_ok=True)
        raise

    resource_entries.extend(quarantines)
    resource_entries.sort(
        key=lambda resource: (
            resource["period"]["year"],
            resource["period"].get("quarter") or "",
            resource["decision"],
        )
    )
    reconciliation = {
        "schema_version": "0.3.0",
        "normalizer_version": NORMALIZER_VERSION,
        "inputs": {
            "catalogue": display_path(catalogue_path),
            "census": display_path(census_path),
            "raw_manifest": display_path(raw_manifest_path),
        },
        "output": {
            "path": display_path(output_path),
            "sha256": file_sha256(output_path),
            "bytes": output_path.stat().st_size,
            "records": total_records,
        },
        "totals": {
            "catalogued_resources": len(plans) + len(quarantines),
            "canonical_resources": len(plans),
            "recovered_resources": sum(
                plan.disposition == "recovered_official_alternate" for plan in plans
            ),
            "bilingual_measure_override_resources": sum(
                plan.disposition == "normalized_with_bilingual_measure_override"
                for plan in plans
            ),
            "quarantined_resources": len(quarantines),
            "canonical_records": total_records,
            "by_decision": totals_by_decision,
        },
        "checks": [
            "Catalogue, census, and raw-manifest English resource UUID sets match.",
            "Every selected source checksum was verified before parsing.",
            "Every accepted schema and decision-measure relationship was validated.",
            "Every canonical record ID is unique.",
            "Any quarantined resource produces no canonical records.",
            "The 2022 Q2 English positive measure override is accepted only while its "
            "paired official French release matches at normalized-record level.",
        ],
        "census_row_note": (
            "The census estimate counts nonempty rows after the first detected header. "
            "In early province-grouped CSVs that includes repeated headers and province headings; "
            "for recovered 2021 Q2 it describes the quarantined catalogue blob, not the replacement."
        ),
        "resources": resource_entries,
    }
    write_json(reconciliation_path, reconciliation)
    write_json(schema_path, build_schema_document())
    return reconciliation


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--catalogue", type=Path, default=CATALOGUE_PATH)
    parser.add_argument("--census", type=Path, default=CENSUS_PATH)
    parser.add_argument("--raw-manifest", type=Path, default=RAW_MANIFEST_PATH)
    parser.add_argument("--output", type=Path, default=OUTPUT_PATH)
    parser.add_argument("--reconciliation", type=Path, default=RECONCILIATION_PATH)
    parser.add_argument("--schema", type=Path, default=SCHEMA_PATH)
    parser.add_argument("--fail-on-quarantine", action="store_true")
    args = parser.parse_args()
    try:
        reconciliation = build(
            catalogue_path=args.catalogue.resolve(),
            census_path=args.census.resolve(),
            raw_manifest_path=args.raw_manifest.resolve(),
            output_path=args.output.resolve(),
            reconciliation_path=args.reconciliation.resolve(),
            schema_path=args.schema.resolve(),
            fail_on_quarantine=args.fail_on_quarantine,
        )
    except CanonicalBuildError as error:
        parser.exit(1, f"error: {error}\n")
    totals = reconciliation["totals"]
    print(
        f"Built {totals['canonical_records']:,} canonical records from "
        f"{totals['canonical_resources']} resources; "
        f"{totals['quarantined_resources']} resource quarantined."
    )
    print(f"CSV: {reconciliation['output']['path']}")


if __name__ == "__main__":
    main()
