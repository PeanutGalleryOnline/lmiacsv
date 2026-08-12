"""Collect a durable evidence bundle for the mislinked 2021 Q2 negative resource.

The collector archives what the Open Government English resource currently serves,
the corresponding 2020 Q2 file, the recovered ESDC English/French workbooks, resource
page HTML, and catalogue JSON-LD. It then writes independent byte- and row-level
verification receipts. Signed redirect query strings are never persisted.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import html
import io
import json
import tempfile
from datetime import datetime, timezone
from html.parser import HTMLParser
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit, urlunsplit
from urllib.request import Request, urlopen


WORKSPACE_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = (
    WORKSPACE_ROOT / "evidence" / "open-government" / "2021-q2-negative-resource" / "2026-08-11"
)
USER_AGENT = "lmiacsv-evidence-collector/0.1"

DATASET_ID = "f82f66f2-a22b-4511-bccf-e1d74db39ae5"
RESOURCE_IDS = {
    "portal_2020_q2_en": "5bb40b78-46ac-430f-b1b9-b8ed25632bc9",
    "portal_2021_q2_en": "4f9a23d3-c6a6-4472-94ad-ec9646351539",
    "portal_2021_q2_fr": "f7fbb221-758d-4a5e-ab61-967d9fb19ddb",
    "portal_2020_q1_en": "93fed83b-266b-45ba-aa65-4f356349f9e7",
    "portal_2020_q3_en": "39f46bb3-f9cd-4111-a2d6-dd64c8381902",
    "portal_2021_q3_en": "f731e9e9-273d-486a-8551-10fbe65be5f5",
}
PORTAL_CSV_FILENAME = (
    "useb-dgcetfw-tetdip-piddiviaionline-publicationemployer-list2020-"
    "employer-list2020q22020q2csv202.csv"
)


def resource_page(resource_id: str) -> str:
    return (
        f"https://open.canada.ca/data/en/dataset/{DATASET_ID}/resource/"
        f"{resource_id}?wbdisable=false"
    )


def portal_download(resource_id: str) -> str:
    return (
        f"https://open.canada.ca/data/dataset/{DATASET_ID}/resource/{resource_id}/"
        f"download/{PORTAL_CSV_FILENAME}"
    )


ASSETS = (
    {
        "id": "portal-2020-q2-en-csv",
        "role": "control: official 2020 Q2 English portal download",
        "url": portal_download(RESOURCE_IDS["portal_2020_q2_en"]),
        "path": f"sources/{PORTAL_CSV_FILENAME.replace('.csv', '-2020-resource.csv')}",
    },
    {
        "id": "portal-2021-q2-en-csv",
        "role": "anomaly: official 2021 Q2 English portal download",
        "url": portal_download(RESOURCE_IDS["portal_2021_q2_en"]),
        "path": f"sources/{PORTAL_CSV_FILENAME.replace('.csv', '-2021-resource.csv')}",
    },
    {
        "id": "esdc-2021-q2-negative-en-xlsx",
        "role": "recovered analytical source: official ESDC English workbook",
        "url": "https://www.esdc.gc.ca/ouvert-open/bca-seb/imt-lmi/TFWP_2021Q2_Negative_EN.xlsx",
        "path": "sources/TFWP_2021Q2_Negative_EN.xlsx",
    },
    {
        "id": "esdc-2021-q2-negative-fr-xlsx",
        "role": "catalogued cross-check: official ESDC French workbook",
        "url": "https://www.esdc.gc.ca/ouvert-open/bca-seb/imt-lmi/TFWP_2021Q2_Negative_FR.xlsx",
        "path": "sources/TFWP_2021Q2_Negative_FR.xlsx",
    },
    *(
        {
            "id": f"{name}-resource-page",
            "role": f"resource-page snapshot: {name.replace('_', ' ')}",
            "url": resource_page(resource_id),
            "path": f"pages/{name.replace('_', '-')}.html",
        }
        for name, resource_id in RESOURCE_IDS.items()
    ),
    {
        "id": "negative-employer-list-catalogue-jsonld",
        "role": "machine-readable Open Government catalogue snapshot",
        "url": f"https://open.canada.ca/data/en/dataset/{DATASET_ID}.jsonld",
        "path": f"catalogue/{DATASET_ID}.jsonld",
    },
)

SAFE_RESPONSE_HEADERS = (
    "content-type",
    "content-length",
    "content-disposition",
    "last-modified",
    "etag",
    "date",
    "server",
)


class VisibleText(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self.ignored_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"script", "style", "template"}:
            self.ignored_depth += 1

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style", "template"} and self.ignored_depth:
            self.ignored_depth -= 1

    def handle_data(self, data: str) -> None:
        if not self.ignored_depth:
            value = " ".join(html.unescape(data).split())
            if value:
                self.parts.append(value)


def sanitize_url(value: str) -> str:
    parsed = urlsplit(value)
    return urlunsplit((parsed.scheme, parsed.netloc, parsed.path, "", ""))


def sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def fetch_asset(asset: dict[str, str], root: Path, retrieved_at: str) -> dict[str, Any]:
    request = Request(asset["url"], headers={"User-Agent": USER_AGENT})
    with urlopen(request, timeout=60) as response:
        body = response.read()
        headers = {
            key: response.headers.get(key)
            for key in SAFE_RESPONSE_HEADERS
            if response.headers.get(key) is not None
        }
        receipt = {
            "id": asset["id"],
            "role": asset["role"],
            "retrieved_at": retrieved_at,
            "requested_url": sanitize_url(asset["url"]),
            "final_url": sanitize_url(response.geturl()),
            "http_status": response.getcode(),
            "response_headers": headers,
            "archive_path": asset["path"],
            "bytes": len(body),
            "sha256": sha256_bytes(body),
        }
    target = root / asset["path"]
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(body)
    return receipt


def validation_status(page: bytes) -> dict[str, str | None]:
    parser = VisibleText()
    parser.feed(page.decode("utf-8", errors="replace"))
    result: dict[str, str | None] = {"status": None, "timestamp": None}
    for index, value in enumerate(parser.parts):
        if value == "Validation status" and index + 1 < len(parser.parts):
            result["status"] = parser.parts[index + 1]
        if value == "Validation timestamp" and index + 1 < len(parser.parts):
            result["timestamp"] = parser.parts[index + 1]
    return result


def inspect_csv(path: Path) -> dict[str, Any]:
    body = path.read_bytes()
    try:
        text = body.decode("utf-8-sig", errors="strict")
        encoding = "utf-8-sig"
    except UnicodeDecodeError:
        text = body.decode("cp1252", errors="strict")
        encoding = "cp1252"
    rows = list(csv.reader(io.StringIO(text)))
    note_index = next(
        (index for index, row in enumerate(rows) if row and row[0].strip() == "Notes:"),
        len(rows),
    )
    data_rows = [row for row in rows[2:note_index] if any(cell.strip() for cell in row)]
    return {
        "encoding": encoding,
        "title": rows[0][0].strip(),
        "headers": [cell.strip() for cell in rows[1]],
        "data_rows": len(data_rows),
        "requested_positions": sum(int(row[5]) for row in data_rows),
        "distinct_employer_strings": len({row[2].strip() for row in data_rows}),
    }


def inspect_workbook(path: Path) -> tuple[dict[str, Any], list[tuple[Any, ...]]]:
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError as error:
        raise RuntimeError("openpyxl 3.1.5 is required to verify the receipts") from error

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if len(workbook.sheetnames) != 1:
            raise ValueError(f"Expected one sheet in {path.name}, found {workbook.sheetnames}")
        sheet = workbook.active
        title = str(sheet.cell(1, 1).value).strip()
        headers = [str(sheet.cell(2, column).value).strip() for column in range(1, 7)]
        note_row = next(
            row_number
            for row_number in range(3, sheet.max_row + 1)
            if str(sheet.cell(row_number, 1).value).strip() == "Notes:"
        )
        rows = [
            tuple(values)
            for values in sheet.iter_rows(
                min_row=3,
                max_row=note_row - 1,
                min_col=1,
                max_col=6,
                values_only=True,
            )
            if any(value is not None and str(value).strip() for value in values)
        ]
        receipt = {
            "sheet": sheet.title,
            "title": title,
            "headers": headers,
            "data_rows": len(rows),
            "first_data_row": 3,
            "last_data_row": note_row - 1,
            "notes_row": note_row,
            "requested_positions": sum(int(row[5]) for row in rows),
            "distinct_employer_strings": len({str(row[2]).strip() for row in rows}),
        }
    finally:
        workbook.close()
    return receipt, rows


def build_verification(root: Path, receipts: list[dict[str, Any]]) -> dict[str, Any]:
    receipt_by_id = {receipt["id"]: receipt for receipt in receipts}
    portal_2020_path = root / receipt_by_id["portal-2020-q2-en-csv"]["archive_path"]
    portal_2021_path = root / receipt_by_id["portal-2021-q2-en-csv"]["archive_path"]
    english_path = root / receipt_by_id["esdc-2021-q2-negative-en-xlsx"]["archive_path"]
    french_path = root / receipt_by_id["esdc-2021-q2-negative-fr-xlsx"]["archive_path"]

    csv_2020 = inspect_csv(portal_2020_path)
    csv_2021 = inspect_csv(portal_2021_path)
    english, english_rows = inspect_workbook(english_path)
    french, french_rows = inspect_workbook(french_path)
    if len(english_rows) != len(french_rows):
        raise ValueError("Recovered English and French workbook row counts differ")

    bilingual_mismatches = []
    for row_number, (en_row, fr_row) in enumerate(zip(english_rows, french_rows), start=3):
        en_noc = str(en_row[4]).split("-", 1)[0].strip()
        fr_noc = str(fr_row[4]).split("-", 1)[0].strip()
        en_identity = (en_row[2], en_row[3], en_noc, int(en_row[5]))
        fr_identity = (fr_row[2], fr_row[3], fr_noc, int(fr_row[5]))
        if en_identity != fr_identity:
            bilingual_mismatches.append(row_number)

    page_statuses = {}
    for receipt in receipts:
        if receipt["id"].endswith("resource-page"):
            page_statuses[receipt["id"]] = validation_status(
                (root / receipt["archive_path"]).read_bytes()
            )

    portal_identical = portal_2020_path.read_bytes() == portal_2021_path.read_bytes()
    verification = {
        "schema_version": "0.1.0",
        "finding": "The English 2021 Q2 Open Government resource serves the 2020 Q2 CSV, while matching English and French 2021 Q2 XLSX workbooks remain available on ESDC infrastructure.",
        "portal_duplicate": {
            "byte_for_byte_identical": portal_identical,
            "sha256_matches": (
                receipt_by_id["portal-2020-q2-en-csv"]["sha256"]
                == receipt_by_id["portal-2021-q2-en-csv"]["sha256"]
            ),
            "portal_2020_q2": csv_2020,
            "portal_2021_q2": csv_2021,
        },
        "recovered_workbooks": {
            "english": english,
            "french": french,
            "language_independent_fields_compared": [
                "employer",
                "address",
                "NOC code",
                "requested positions",
            ],
            "row_level_mismatch_count": len(bilingual_mismatches),
            "row_level_mismatch_source_rows": bilingual_mismatches,
        },
        "resource_page_validation": page_statuses,
        "checks": {
            "portal_files_identical": portal_identical,
            "portal_files_describe_2020_q2": (
                "2020" in csv_2020["title"] and "2020" in csv_2021["title"]
            ),
            "recovered_english_describes_2021_q2": ("April to June 2021" in english["title"]),
            "recovered_french_describes_2021_q2": ("avril à juin 2021" in french["title"]),
            "recovered_language_row_counts_match": english["data_rows"] == french["data_rows"],
            "recovered_language_position_totals_match": (
                english["requested_positions"] == french["requested_positions"]
            ),
            "recovered_language_independent_fields_match": not bilingual_mismatches,
        },
    }
    if not all(verification["checks"].values()):
        raise ValueError(f"Evidence verification failed: {verification['checks']}")
    return verification


def write_readme(root: Path, retrieved_at: str, verification: dict[str, Any]) -> None:
    english = verification["recovered_workbooks"]["english"]
    content = f"""# 2021 Q2 negative LMIA resource receipts

Collected at: `{retrieved_at}`

This immutable bundle preserves the evidence for a linkage error in the Open
Government English 2021 Q2 negative LMIA employer resource.

## Finding

The English 2021 Q2 resource serves a byte-for-byte copy of the 2020 Q2 CSV. The
downloaded CSV's own title identifies April–June 2020. Matching English and French
2021 Q2 XLSX workbooks remain available on official ESDC infrastructure.

The recovered workbooks contain {english["data_rows"]:,} data rows and
{english["requested_positions"]:,} requested positions. Employer names, addresses,
NOC codes, and requested-position counts match between the English and French
workbooks on every row.

## Contents

- `sources/`: the two portal CSV downloads and two recovered ESDC workbooks.
- `pages/`: dated HTML snapshots of the affected and comparison resource pages.
- `catalogue/`: the machine-readable negative-employer-list JSON-LD snapshot.
- `receipts.json`: requested/final URLs, selected response headers, sizes, and hashes.
- `verification.json`: independent CSV, workbook, duplicate, and bilingual checks.
- `SHA256SUMS`: checksums for every downloaded source asset.

Final URLs and response headers are sanitized so temporary signed-query parameters
are not persisted. The recovered English URL was inferred from the catalogued French
workbook's official ESDC URL and then validated against that French workbook.
"""
    (root / "README.md").write_text(content, encoding="utf-8")


def collect(output: Path) -> None:
    if output.exists():
        raise FileExistsError(f"Refusing to overwrite immutable evidence bundle: {output}")
    output.parent.mkdir(parents=True, exist_ok=True)
    retrieved_at = datetime.now(timezone.utc).replace(microsecond=0).isoformat()
    with tempfile.TemporaryDirectory(prefix=f".{output.name}-", dir=output.parent) as temporary:
        staging = Path(temporary)
        receipts = [fetch_asset(asset, staging, retrieved_at) for asset in ASSETS]
        verification = build_verification(staging, receipts)
        (staging / "receipts.json").write_text(
            json.dumps(
                {
                    "schema_version": "0.1.0",
                    "retrieved_at": retrieved_at,
                    "dataset_id": DATASET_ID,
                    "assets": receipts,
                },
                indent=2,
                ensure_ascii=False,
            )
            + "\n",
            encoding="utf-8",
        )
        (staging / "verification.json").write_text(
            json.dumps(verification, indent=2, ensure_ascii=False) + "\n",
            encoding="utf-8",
        )
        checksum_rows = sorted((receipt["sha256"], receipt["archive_path"]) for receipt in receipts)
        (staging / "SHA256SUMS").write_text(
            "".join(f"{digest}  {path}\n" for digest, path in checksum_rows),
            encoding="utf-8",
        )
        write_readme(staging, retrieved_at, verification)
        staging.rename(output)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    output = args.output.resolve()
    collect(output)
    verification = json.loads((output / "verification.json").read_text(encoding="utf-8"))
    recovered = verification["recovered_workbooks"]["english"]
    print(f"Evidence bundle written to {output}")
    print(
        f"Verified portal duplicate and recovered {recovered['data_rows']:,} English rows "
        f"against the French workbook with zero mismatches."
    )


if __name__ == "__main__":
    main()
